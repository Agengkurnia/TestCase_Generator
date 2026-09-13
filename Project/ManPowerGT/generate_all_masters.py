"""Workbook test case: semua master PowerGT (Man Power GT)."""
from pathlib import Path

from generate_product_sample import (
    CASES as PRODUCT_CASES,
    CELL_FONT,
    GREEN,
    MOD as MOD_PRODUCT,
    SIT_HEADERS,
    SIT_ROWS as PRODUCT_SIT,
    SKN_HEADERS,
    TITLE_FONT,
    WRAP,
    autosize,
    style_header,
    write_rows,
)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent / "ManPowerGT_Master_TestCase.xlsx"

EMPTY = (
    "Tampil ~/assets/img/data_not_found.png (tinggi ~280px) dan teks "
    "No data available to show. Blok empty mengisi area scroll (~360px)."
)
COL_SEARCH = (
    "1. Ketik di input Search, jangan Enter/klik tombol. Grid tidak berubah.\n"
    "2. Klik tombol search hijau (atau Enter). Grid terfilter.\n"
    "3. Kosongkan input, Enter: filter kolom itu dilepas."
)


def p(menu, extra=""):
    base = (
        f"User login MAVEN, akses menu {menu} (bitView). "
        f"Buka Master Data / {menu}."
    )
    return base + ((" " + extra) if extra else "")


def pe(menu, extra=""):
    return p(menu, "Role punya bitEdit atau SuperUser." + ((" " + extra) if extra else ""))


def pv(menu):
    return (
        f"User login bitView saja (tanpa bitEdit/SuperUser). "
        f"Buka Master Data / {menu}."
    )


REASON = [
    (
        "Master Reason",
        "Index - Page Load",
        "Positive",
        "Halaman Reason: 6 kartu statistik, DataTable, Add sesuai role.",
        pe("Reason"),
        "1. Buka Master Data / Reason.\n2. Tunggu skeleton kartu dan processing grid selesai.",
        "Judul: Reason.\n"
        "Kartu: Total, Active, Inactive, Return, Visit, Order. Total Filtered.\n"
        "Kolom: No, Reason Name, Description, Type, Status. Tidak ada kolom Action.\n"
        "Nama berupa tautan hijau tebal. Type berupa badge berwarna.\n"
        "Inactive memakai badge merah (bg-danger). Add Reason tampil.",
        "High",
    ),
    (
        "Master Reason",
        "Index - Empty State",
        "Positive",
        "Grid kosong menampilkan empty state.",
        p("Reason", "Tidak ada data, atau filter 0 row."),
        "1. Buka Reason atau terapkan filter tanpa hasil.",
        EMPTY,
        "Medium",
    ),
    (
        "Master Reason",
        "Add Reason - Role",
        "Positive",
        "Add Reason hanya untuk bitEdit/SuperUser.",
        pe("Reason"),
        "1. Buka Reason sebagai user edit.",
        "Tombol Add Reason tampil.",
        "High",
    ),
    (
        "Master Reason",
        "Filter Stat Cards - Status",
        "Positive",
        "Kartu Total / Active / Inactive memfilter grid.",
        p("Reason", "Ada data Active dan Inactive."),
        "1. Klik Active, lalu Inactive, lalu Total.",
        "Hanya satu kartu Filtered.\n"
        "Active: status Active. Inactive: status Inactive. Total: semua.\n"
        "Search kolom dan select Type/Status di header dikosongkan saat ganti kartu.\n"
        "Hitungan dari GetList: intTotalReason, intTotalAktif, Inactive = Total-Aktif.",
        "High",
    ),
    (
        "Master Reason",
        "Filter Stat Cards - Type",
        "Positive",
        "Kartu Return / Visit / Order memfilter Type. Visit memakai nilai DB Kunjungan (SIT-010 / SIT-012).",
        p("Reason", "Ada data Return, Kunjungan, Order."),
        "1. Klik kartu Return. Amati Type di grid.\n"
        "2. Klik kartu Visit.\n"
        "3. Klik kartu Order.\n"
        "4. Bandingkan hitungan kartu Visit dengan intTotalKunjungan.",
        "Return: hanya Type Return.\n"
        "Visit: filter txtTipe=Kunjungan (bukan string Visit). Badge di grid tetap label Visit.\n"
        "Order: Type Order.\n"
        "Select filter Type di header tersinkron (nilai Kunjungan saat kartu Visit).\n"
        "Hitungan: cntReturn, cntVisit (intTotalKunjungan), cntOrder.",
        "High",
    ),
    (
        "Master Reason",
        "Column Search dan Select Type/Status",
        "Positive",
        "Name/Description: tombol search. Type/Status: select (change langsung).",
        p("Reason", "Beberapa reason berbeda nama dan type."),
        COL_SEARCH + "\n"
        "4. Pada kolom Type pilih satu type (nilai ParamDetailName, mis. Kunjungan).\n"
        "5. Pada kolom Status pilih Active lalu Inactive.",
        "Name/Description tidak filter on keyup.\n"
        "Select Type memfilter txtTipe (nilai DB, bukan label Visit).\n"
        "Select Status: true=Active, false=Inactive.",
        "High",
    ),
    (
        "Master Reason",
        "Type Badge Warna dan Label",
        "Positive",
        "Tiap Type beda warna; Kunjungan tampil sebagai Visit (SIT-008 / SIT-012).",
        p("Reason", "Ada Type Return, Kunjungan, Order, dan type lain dari ParamApi."),
        "1. Amati kolom Type di list untuk tiap type.\n"
        "2. Bandingkan nilai di DB vs label UI.",
        "Return->Return, Kunjungan->Visit, Order->Order. Type lain tampil apa adanya.\n"
        "Warna badge dari palet primary/warning/success/info/danger/dark, satu warna per type.\n"
        "Nilai yang disimpan/dikirim ke API tetap ParamDetailName (Kunjungan, bukan Visit).",
        "High",
    ),
    (
        "Master Reason",
        "Open Detail",
        "Positive",
        "Klik nama membuka form in-page.",
        p("Reason", "Minimal 1 reason."),
        "1. Klik tautan Reason Name.",
        "List hilang, form tampil. Back + Save Reason (Save jika edit).\n"
        "Judul Edit Reason. Field terisi GetByGuid. Tidak pindah halaman Detail.",
        "High",
    ),
    (
        "Master Reason",
        "Type setelah Create",
        "Positive",
        "Setelah create, Type di form/list sesuai DB (SIT-006). Perhatikan mapping lowercase di renderForm.",
        pe("Reason"),
        "1. Add Reason. Isi Name, pilih Type (contoh Kunjungan), Description opsional, Status Active. Save.\n"
        "2. Buka kembali detail reason yang baru.\n"
        "3. Cek Type di form dan di list.",
        "Save sukses: Reason created. Type di list = badge Visit jika DB Kunjungan.\n"
        "Form detail harus menampilkan Type yang sama (dropdown terpilih).\n"
        "Jika dropdown kosong padahal DB ada txtTipe: Fail "
        "(renderForm memakai val(tipe.toLowerCase()) vs option Case-sensitive).",
        "High",
    ),
    (
        "Master Reason",
        "Type Visit / Others",
        "Positive",
        "Type dari ParamApi REASON_CATEGORY_GT yang IsActive boleh disimpan, termasuk Visit/Kunjungan dan Others/Lainnya (SIT-007).",
        pe("Reason", "ParamApi mengembalikan Kunjungan dan/atau Lainnya (atau nama setara) yang IsActive."),
        "1. Add Reason.\n"
        "2. Pilih Type Visit/Kunjungan, Save.\n"
        "3. Ulangi Add dengan Type Others/Lainnya jika ada di dropdown.",
        "Tidak ada error tidak enable input data dengan Tipe Reason Visit/Others.\n"
        "Data tersimpan. Type di list sesuai mapping label.\n"
        "Hanya opsi IsActive dari REASON_CATEGORY_GT yang muncul.",
        "High",
    ),
    (
        "Master Reason",
        "Reason Name Mandatory",
        "Negative",
        "Name wajib.",
        pe("Reason"),
        "1. Add Reason. Kosongkan Name. Pilih Type. Save.",
        "Pesan: Reason Name is required. Tidak tersimpan.",
        "High",
    ),
    (
        "Master Reason",
        "Type Mandatory",
        "Negative",
        "Type wajib.",
        pe("Reason"),
        "1. Add Reason. Isi Name. Type kosong. Save.",
        "Pesan: Type is required. Tidak tersimpan.",
        "High",
    ),
    (
        "Master Reason",
        "Save Create Update",
        "Positive",
        "Create POST Save, edit POST Update/{guid}, konfirmasi Save this reason?",
        pe("Reason"),
        "1. Add, isi wajib, Save, batalkan konfirmasi lalu konfirmasi ya.\n"
        "2. Buka detail, ubah Description, Save.",
        "Dialog: Save this reason?\n"
        "Create: Reason created. Edit: Reason updated.\n"
        "Kembali ke list, grid reload.",
        "High",
    ),
    (
        "Master Reason",
        "Confirm Back",
        "Positive",
        "Back minta konfirmasi.",
        pe("Reason", "Form terbuka."),
        "1. Klik Back. Batalkan. Klik Back lagi, ya.",
        "Dialog: Back to reason list? Batal tetap di form. Ya kembali list.",
        "Medium",
    ),
    (
        "Master Reason",
        "Delete Unwired",
        "Negative",
        "Tidak ada Delete di API; tombol Delete disembunyikan.",
        pe("Reason", "Form terbuka."),
        "1. Amati header detail. Cari Delete.",
        "btnDeleteReason class d-none, tidak ter-bind. Tidak ada endpoint Delete.",
        "Medium",
    ),
    (
        "Master Reason",
        "Inserted / Updated By",
        "Positive",
        "User tercatat di DB setelah create/update (SIT-009). Payload FE tidak mengirim InsertedBy.",
        pe("Reason"),
        "1. Create Reason, Save.\n"
        "2. Cek kolom Inserted By di mGtReason.\n"
        "3. Update, cek Updated By.",
        "Kolom terisi nama/user login, bukan GUID mentah.\n"
        "Jika kosong/GUID: Fail (backend).",
        "High",
    ),
]

CHANNEL = [
    (
        "Master Channel",
        "Index - View Only",
        "Positive",
        "Channel list-only, badge View-only, tanpa form detail.",
        p("Channel"),
        "1. Buka Master Data / Channel.\n2. Amati header dan grid.",
        "Judul Channel. Badge View-only.\n"
        "Tidak ada Add / Save / Back detail.\n"
        "Kolom: No, Channel, Type Customer, Account.\n"
        "Filter: tiga select (bukan input+tombol search).\n"
        "Klik baris tidak membuka detail.",
        "High",
    ),
    (
        "Master Channel",
        "Load Join Outbound",
        "Positive",
        "Grid hasil join OutletTypePartner (IsActive) + OutletAccountType by OutletAccountTypeId.",
        p("Channel", "Outbound Master Data API hidup."),
        "1. Buka Channel, tunggu skeleton hilang.",
        "Channel = OutletChannelTypeName dari account type.\n"
        "Type Customer = TypePartnerName. Account = OutletAccountTypeName.\n"
        "Hanya TypePartner IsActive. Gagal API: Failed to load Channel list, grid kosong.",
        "High",
    ),
    (
        "Master Channel",
        "Column Filter Select",
        "Positive",
        "Select Channel / Type Customer / Account filter exact match.",
        p("Channel", "Ada lebih dari satu nilai per kolom."),
        "1. Pilih satu Channel di select kolom 1.\n"
        "2. Tambahkan filter Type Customer.\n"
        "3. Kembalikan ke All Channel.",
        "Filter exact (regex ^nilai$). Kombinasi AND antar kolom.\n"
        "Opsi select terisi unik dari data, terurut.",
        "High",
    ),
    (
        "Master Channel",
        "Empty State",
        "Positive",
        "Empty state Channel memakai copy berbeda.",
        p("Channel", "API gagal atau hasil 0."),
        "1. Kondisi tanpa data (atau gagal load).",
        "Gambar data_not_found.png. Teks: no data available in table (huruf kecil, beda dari master lain).",
        "Medium",
    ),
]

CUSTOMER = [
    (
        "Master Customer",
        "Index - Page Load",
        "Positive",
        "Customer read-only list: 4 kartu termasuk With GPS.",
        p("Customer"),
        "1. Buka Master Data / Customer.",
        "Kartu: Total, Active, Inactive, With GPS. Total Filtered.\n"
        "Kolom: No, Code, Name, City, Channel, Salesman, Status.\n"
        "Name tautan hijau. Tidak ada tombol Add.",
        "High",
    ),
    (
        "Master Customer",
        "Filter With GPS",
        "Positive",
        "Kartu With GPS memfilter bitHasGps.",
        p("Customer", "Ada customer dengan dan tanpa GPS."),
        "1. Klik With GPS. 2. Klik Total.",
        "Hanya baris dengan GPS. Kartu With GPS Filtered. Search kolom ter-reset.",
        "High",
    ),
    (
        "Master Customer",
        "Column Search",
        "Positive",
        "Search per kolom klik/Enter, bukan keyup.",
        p("Customer"),
        COL_SEARCH,
        "Grid terfilter setelah klik/Enter saja.",
        "High",
    ),
    (
        "Master Customer",
        "Open Detail",
        "Positive",
        "Klik Name membuka detail in-page read-only.",
        p("Customer", "Minimal 1 customer."),
        "1. Klik nama customer.",
        "List hilang. Header Back saja (tidak ada Save).\n"
        "Profil: nama, kode, status, channel.\n"
        "Section: Store & Contact, Address, Location Map, Store Photo.\n"
        "Semua input disabled. GetByGuid mengisi sisa field.",
        "High",
    ),
    (
        "Master Customer",
        "GPS dan Peta",
        "Positive",
        "Koordinat valid tampil di GPS + peta + Google Maps. 0,0 dianggap invalid.",
        p("Customer", "Satu customer dengan lat/lng valid, satu tanpa/0,0."),
        "1. Buka detail yang punya GPS valid. Amati GPS, tombol Maps, peta.\n"
        "2. Buka detail tanpa koordinat valid.",
        "Valid: GPS lat, lng. Tombol Google Maps href maps?q=lat,lng. Marker Leaflet.\n"
        "Invalid (kosong, NaN, 0,0, di luar range): GPS —, tombol Maps hidden, "
        "teks No valid coordinates to show on the map.",
        "High",
    ),
    (
        "Master Customer",
        "Store Photo",
        "Positive",
        "Klik thumbnail foto toko membuka modal Store Photo.",
        p("Customer", "Customer punya foto."),
        "1. Buka detail. Klik foto jika ada.",
        "Modal Store Photo menampilkan gambar. Tanpa foto: placeholder, klik tidak wajib buka modal.",
        "Medium",
    ),
    (
        "Master Customer",
        "Confirm Back",
        "Positive",
        "Back konfirmasi.",
        p("Customer", "Detail terbuka."),
        "1. Klik Back.",
        "Dialog: Back to customer list?",
        "Medium",
    ),
    (
        "Master Customer",
        "Empty State",
        "Positive",
        "Grid kosong.",
        p("Customer", "0 row."),
        "1. Kondisi tanpa data.",
        EMPTY,
        "Medium",
    ),
]

EMPLOYEE = [
    (
        "Master Employee",
        "Index - Page Load",
        "Positive",
        "Employee list: kartu Total / Motoris / SPG GT / Active. CSV toolbar.",
        pe("Employee"),
        "1. Buka Master Data / Employee.",
        "Judul Employee. Kartu Total Filtered, Motoris, SPG GT, Active.\n"
        "Kolom: No, NIK, Name, Role, Branch, Region, Status. Name tautan hijau.\n"
        "Download Data tampil. Upload Data + Edit tampil (bitEdit).",
        "High",
    ),
    (
        "Master Employee",
        "Filter Role Cards",
        "Positive",
        "Kartu Motoris / SPG GT / Active memfilter grid.",
        p("Employee", "Ada Motoris dan SPG GT."),
        "1. Klik Motoris. 2. SPG GT. 3. Active. 4. Total.",
        "Motoris: cardFilter motoris. SPG GT: spg gt. Active: active.\n"
        "Hitungan: intTotalPegawai, intTotalMotoris, intTotalSpgGt, intTotalAktif.",
        "High",
    ),
    (
        "Master Employee",
        "Column Search",
        "Positive",
        "Search per kolom klik/Enter.",
        p("Employee"),
        COL_SEARCH,
        "Tidak auto-filter keyup.",
        "High",
    ),
    (
        "Master Employee",
        "Open Detail",
        "Positive",
        "Detail read-only: Profile, Assignment, status history.",
        p("Employee", "Minimal 1 pegawai."),
        "1. Klik nama. 2. Amati field dan tabel history.",
        "Header Back. Profile: NIK, Name, Role, Phone, Status (disabled).\n"
        "Assignment: Branch, Region, Notes.\n"
        "Status history terisi GetByGuid. Tidak ada Save.",
        "High",
    ),
    (
        "Master Employee",
        "Download Data",
        "Positive",
        "Download Data mengarah ke ExportExcel.",
        p("Employee"),
        "1. Klik Download Data.",
        "Request GtEmployee/ExportExcel. File terunduh / tidak 403 untuk user berhak.",
        "Medium",
    ),
    (
        "Master Employee",
        "Upload CSV",
        "Positive",
        "Upload Data menerima .csv, POST FormData file.",
        pe("Employee", "Siapkan CSV valid."),
        "1. Klik Upload Data, pilih .csv valid.",
        "accept=.csv. Sukses: Upload successful, grid+summary reload. Input file dikosongkan.\n"
        "Gagal: swal error, input dikosongkan.",
        "High",
    ),
    (
        "Master Employee",
        "Edit Spreadsheet Modal",
        "Positive",
        "Edit membuka modal grid, Save mengirim CSV pegawai_edit.csv.",
        pe("Employee"),
        "1. Klik Edit. 2. Ubah satu sel. 3. Save di footer modal (konfirmasi jika ada).",
        "Modal Edit. Kolom: NIK, Name, Role, Branch, Region, Phone, Notes, Status.\n"
        "Save memanggil Upload. Grid list reload.",
        "High",
    ),
    (
        "Master Employee",
        "Upload Hidden View Only",
        "Negative",
        "View-only tidak melihat Upload/Edit.",
        pv("Employee"),
        "1. Buka Employee.",
        "Upload Data dan Edit tidak dirender. Download Data boleh tetap ada.",
        "High",
    ),
    (
        "Master Employee",
        "Confirm Back",
        "Positive",
        "Back dari detail.",
        p("Employee", "Detail terbuka."),
        "1. Klik Back.",
        "Dialog: Back to employee list?",
        "Medium",
    ),
    (
        "Master Employee",
        "Empty State",
        "Positive",
        "Grid kosong.",
        p("Employee", "0 row."),
        "1. Kondisi tanpa data.",
        EMPTY,
        "Medium",
    ),
]

STOKIS = [
    (
        "Master Stokis",
        "Index - Page Load",
        "Positive",
        "Stokis list + CSV toolbar + kartu With GPS.",
        pe("Stokis"),
        "1. Buka Master Data / Stokis.",
        "Kartu Total, Active, Inactive, With GPS.\n"
        "Kolom: Outlet ID, Name, Branch, Region, Status (plus No). Name tautan hijau.\n"
        "Download Data. Upload+Edit jika bitEdit.",
        "High",
    ),
    (
        "Master Stokis",
        "Filter With GPS",
        "Positive",
        "Kartu GPS memfilter stokis yang punya koordinat.",
        p("Stokis", "Ada stokis dengan dan tanpa lat/lng."),
        "1. Klik With GPS.",
        "Hanya row dengan GPS. cardFilter=gps.",
        "High",
    ),
    (
        "Master Stokis",
        "Column Search",
        "Positive",
        "Search per kolom klik/Enter.",
        p("Stokis"),
        COL_SEARCH,
        "Tidak auto-filter keyup.",
        "High",
    ),
    (
        "Master Stokis",
        "Open Detail",
        "Positive",
        "Detail read-only: Profile, Location, peta, stock nested.",
        p("Stokis", "Minimal 1 stokis."),
        "1. Klik nama stokis. 2. Scroll semua section.",
        "Header Back. Profile: Outlet ID, Name, Phone, Status.\n"
        "Location: Branch, Region, Address, City, GPS.\n"
        "Map, Stock Summary (Total SKU, Qty Carton, Last Updated),\n"
        "Stock by Product, Stock Input History, Status History.\n"
        "GetByGuid mengisi nested. Semua field disabled.",
        "High",
    ),
    (
        "Master Stokis",
        "GPS dan Peta",
        "Positive",
        "Koordinat valid: GPS + Google Maps + Leaflet. Invalid: empty map.",
        p("Stokis"),
        "1. Detail dengan GPS valid. 2. Detail tanpa GPS / 0,0.",
        "Valid: lat, lng, tombol Maps, marker.\n"
        "Invalid: GPS —, Maps hidden, No valid coordinates to show on the map.",
        "High",
    ),
    (
        "Master Stokis",
        "Nested Stock Tables",
        "Positive",
        "Sub-grid stock dan history empty state sama.",
        p("Stokis", "Satu stokis dengan stock, satu tanpa."),
        "1. Buka stokis ber-stock. 2. Buka stokis tanpa stock.",
        "Ada data: Code, Product, Qty terisi. History punya Date/Motoris/Notes.\n"
        "Kosong: empty state data_not_found + No data available to show.",
        "Medium",
    ),
    (
        "Master Stokis",
        "Download Data",
        "Positive",
        "Download ke ExportExcel.",
        p("Stokis"),
        "1. Klik Download Data.",
        "Request GtStockist/ExportExcel.",
        "Medium",
    ),
    (
        "Master Stokis",
        "Upload CSV",
        "Positive",
        "Upload .csv. Error bisa menampilkan daftar errors.",
        pe("Stokis"),
        "1. Upload CSV valid. 2. Upload CSV invalid.",
        "Sukses: Upload successful, reload.\n"
        "Gagal: pesan + hingga 5 baris errors. Input dikosongkan. accept=.csv.",
        "High",
    ),
    (
        "Master Stokis",
        "Edit Spreadsheet Modal",
        "Positive",
        "Edit modal kolom termasuk lat/lng.",
        pe("Stokis"),
        "1. Klik Edit. 2. Ubah Name atau Status. Save.",
        "Kolom: Outlet ID, Name, Branch, Region, Phone, Address, Latitude, Longitude, Status.\n"
        "File upload: stokis_edit.csv. List reload.",
        "High",
    ),
    (
        "Master Stokis",
        "Confirm Back",
        "Positive",
        "Back konfirmasi.",
        p("Stokis", "Detail terbuka."),
        "1. Klik Back.",
        "Dialog: Back to stokis list?",
        "Medium",
    ),
    (
        "Master Stokis",
        "Empty State",
        "Positive",
        "Grid kosong.",
        p("Stokis", "0 row."),
        "1. Kondisi tanpa data.",
        EMPTY,
        "Medium",
    ),
]

TAX = [
    (
        "Master Tax",
        "Index - Read Only",
        "Positive",
        "Tax dari GetRatePpn, Type selalu PPN. Add disembunyikan JS.",
        p("Tax"),
        "1. Buka Master Data / Tax. Amati tombol Add dan grid.",
        "Kartu Total/Active/Inactive.\n"
        "Kolom: No, Type (badge PPN), Tax Code (tautan hijau), Tax Name, Percentage (%), DPP Value, Status.\n"
        "btnAddTax di-hide di init meski Razor merender untuk user edit. Tidak ada create.",
        "High",
    ),
    (
        "Master Tax",
        "Filter Stat Cards",
        "Positive",
        "Filter Active/Inactive client-side dari cache.",
        p("Tax", "Ada tax Active dan Inactive."),
        "1. Klik Active, Inactive, Total.",
        "Filter cache GetRatePpn. Satu kartu Filtered. Search kolom ter-reset.",
        "High",
    ),
    (
        "Master Tax",
        "Column Search",
        "Positive",
        "Search per kolom klik/Enter pada cache.",
        p("Tax"),
        COL_SEARCH,
        "Tidak keyup. Termasuk search Status.",
        "High",
    ),
    (
        "Master Tax",
        "Open View Tax",
        "Positive",
        "Detail View Tax, semua field readonly, Save/Delete hidden.",
        p("Tax", "Minimal 1 tax."),
        "1. Klik Tax Code.",
        "Judul View Tax. Type, Code, Name, Percentage, DPP, Effective From/To, Status: disabled.\n"
        "Tanggal format DD Mon YYYY (formatDisplayDate).\n"
        "btnSaveTax dan btnDeleteTax d-none. Hanya Back.",
        "High",
    ),
    (
        "Master Tax",
        "Confirm Back",
        "Positive",
        "Back konfirmasi.",
        p("Tax", "Detail terbuka."),
        "1. Klik Back.",
        "Dialog: Back to tax list?",
        "Medium",
    ),
    (
        "Master Tax",
        "Load Failure",
        "Negative",
        "Gagal GetRatePpn: hitungan 0 + error.",
        p("Tax", "Simulasikan API gagal jika memungkinkan."),
        "1. Buka Tax saat GetRatePpn gagal.",
        "Failed to load tax rates. Kartu 0. Empty state di grid.",
        "Medium",
    ),
    (
        "Master Tax",
        "Empty State",
        "Positive",
        "Grid kosong.",
        p("Tax", "0 row."),
        "1. Tanpa data.",
        EMPTY,
        "Medium",
    ),
]

LIMIT = [
    (
        "Master Limit",
        "Index - Page Load",
        "Positive",
        "Limit list + Add. Kolom Daily Min/Max.",
        pe("Limit"),
        "1. Buka Master Data / Limit.",
        "Kartu Total/Active/Inactive.\n"
        "Kolom: No, Limit Name (tautan hijau), Jabatan, Daily Min, Daily Max, Status.\n"
        "Add Limit tampil untuk edit.",
        "High",
    ),
    (
        "Master Limit",
        "Filter Stat Cards",
        "Positive",
        "Filter Active/Inactive.",
        p("Limit"),
        "1. Klik Active, Inactive, Total.",
        "Satu kartu Filtered. Search kolom ter-reset.",
        "High",
    ),
    (
        "Master Limit",
        "Column Search",
        "Positive",
        "Search Limit Name / Jabatan / Status.",
        p("Limit"),
        COL_SEARCH,
        "Tidak keyup.",
        "High",
    ),
    (
        "Master Limit",
        "Add Form Defaults",
        "Positive",
        "Form Add: history hidden, End Date Ongoing, Start Date min=hari ini.",
        pe("Limit"),
        "1. Klik Add Limit.",
        "Judul Add Limit. Limit Name kosong. Jabatan Select Jabatan...\n"
        "Type Jabatan tergantung Jabatan. Status Active.\n"
        "Daily Min/Max, Weekly/Monthly HKE kosong. End Date Ongoing disabled.\n"
        "Start Date min hari ini, nilai awal kosong. Version history tersembunyi.\n"
        "Jabatan dari PositionApi GetAllByAppCode.",
        "High",
    ),
    (
        "Master Limit",
        "Jabatan Type Depends",
        "Positive",
        "Ganti Jabatan mengisi ulang Type Jabatan dari PositionDetailTypes.",
        pe("Limit", "Minimal 2 jabatan, salah satunya punya type."),
        "1. Add Limit. Pilih Jabatan A, amati Type.\n2. Ganti Jabatan B.",
        "Type Jabatan di-reset lalu diisi type milik jabatan terpilih. Gagal API: Failed to load Jabatan list.",
        "High",
    ),
    (
        "Master Limit",
        "Limit Name Mandatory",
        "Negative",
        "Limit Name wajib.",
        pe("Limit"),
        "1. Add. Pilih Jabatan. Kosongkan Limit Name. Save, konfirmasi ya.",
        "Pesan: Limit Name is required.",
        "High",
    ),
    (
        "Master Limit",
        "Jabatan Mandatory",
        "Negative",
        "Jabatan wajib.",
        pe("Limit"),
        "1. Add. Isi Limit Name. Jabatan kosong. Save, konfirmasi ya.",
        "Pesan: Jabatan is required.",
        "High",
    ),
    (
        "Master Limit",
        "Daily Min Max",
        "Negative",
        "Daily Min tidak boleh > Daily Max.",
        pe("Limit"),
        "1. Add. Isi Name + Jabatan. Daily Min=10, Daily Max=5. Save, konfirmasi ya.",
        "Pesan: Daily Min cannot be greater than Daily Max. Field angka (isNumberKey).",
        "High",
    ),
    (
        "Master Limit",
        "Save Create",
        "Positive",
        "Create: konfirmasi Save this new Limit?",
        pe("Limit"),
        "1. Isi Name, Jabatan, angka valid. Save. Batalkan lalu ya.",
        "Dialog create: Save this new Limit?\nSukses: Limit created, kembali list, summary reload.",
        "High",
    ),
    (
        "Master Limit",
        "Save Update New Version",
        "Positive",
        "Edit selalu versi baru dan menutup versi aktif.",
        pe("Limit", "Ada limit tersimpan."),
        "1. Buka detail. Ubah Daily Max. Save.",
        "Dialog: Saving adds a new version and automatically closes the currently active version(s). Continue?\n"
        "Sukses: New version saved. History menampilkan versi. Versi lama closed, yang baru aktif.",
        "High",
    ),
    (
        "Master Limit",
        "Version History",
        "Positive",
        "Edit menampilkan daftar versi; klik baris mengisi detail versi.",
        pe("Limit", "Limit punya >=1 versi."),
        "1. Buka detail limit existing.\n2. Klik satu baris history.",
        "History tampil. Klik row highlight. Panel: Type Jabatan, Active, Daily Min/Max, Weekly/Monthly HKE.",
        "Medium",
    ),
    (
        "Master Limit",
        "Confirm Back",
        "Positive",
        "Back konfirmasi.",
        pe("Limit", "Form terbuka."),
        "1. Klik Back.",
        "Dialog: Back to limit list? Add tampil lagi di list.",
        "Medium",
    ),
    (
        "Master Limit",
        "Empty State",
        "Positive",
        "Grid kosong.",
        p("Limit", "0 row."),
        "1. Tanpa data.",
        EMPTY,
        "Medium",
    ),
    (
        "Master Limit",
        "View Only No Add",
        "Negative",
        "Tanpa bitEdit: tidak ada Add/Save.",
        pv("Limit"),
        "1. Buka Limit. Buka detail dari nama.",
        "Add Limit dan Save Limit tidak dirender. Detail tetap bisa dilihat (GetByGuid).",
        "High",
    ),
]


def main():
    other = REASON + CHANNEL + CUSTOMER + EMPLOYEE + STOKIS + TAX + LIMIT
    skn = []
    for field, tipe, desc, pre, steps, exp, pri in PRODUCT_CASES:
        skn.append((MOD_PRODUCT, field, tipe, desc, pre, steps, exp, pri))
    skn.extend(other)

    wb = Workbook()
    ws = wb.active
    ws.title = "Test Case (SKN)"
    ws["A1"] = "TEST CASE — MAN POWER GT — MASTER DATA"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(SKN_HEADERS))
    ws["A2"] = (
        "Sumber: MAVEN Views/PowerGT/Master/* dan wwwroot/js/powergt/master/*. "
        "Filter kolom Modul / Area untuk satu master. Status Uji / Hasil Aktual / QA / Tanggal Uji diisi saat eksekusi. "
        "Sheet Issue SIT: log lama Product+Reason, Ref ID = Scenario ID di sheet ini."
    )
    ws["A2"].alignment = WRAP
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(SKN_HEADERS))
    ws.row_dimensions[2].height = 42

    for c, h in enumerate(SKN_HEADERS, 1):
        ws.cell(3, c, h)
    style_header(ws, 3, len(SKN_HEADERS))

    rows = []
    for i, (mod, field, tipe, desc, pre, steps, exp, pri) in enumerate(skn, 1):
        rows.append([
            i, f"SKN-{i:03d}", mod, field, tipe, desc, pre, steps, exp, pri,
            "Not Run", "", "", "",
        ])
    write_rows(ws, 4, rows)
    for r in range(4, 4 + len(rows)):
        ws.row_dimensions[r].height = 72
    ws.auto_filter.ref = f"A3:N{3 + len(rows)}"
    ws.freeze_panes = "A4"
    ws.row_dimensions[3].height = 22
    autosize(ws, [6, 14, 18, 32, 12, 44, 36, 52, 52, 12, 12, 22, 14, 16])

    dv_status = DataValidation(type="list", formula1='"Not Run,Pass,Fail"', allow_blank=True)
    dv_status.add(f"K4:K{3 + len(rows)}")
    ws.add_data_validation(dv_status)
    dv_tipe = DataValidation(type="list", formula1='"Positive,Negative"', allow_blank=False)
    dv_tipe.add(f"E4:E{3 + len(rows)}")
    ws.add_data_validation(dv_tipe)

    # Reason SIT Ref IDs: SKN numbers depend on Product=25 then Reason order
    # Product 1-25, Reason starts 26:
    # 26 Index, 27 Empty, 28 Add role, 29 Filter status, 30 Filter type,
    # 31 Column, 32 Badge, 33 Open, 34 Type after create, 35 Visit/Others,
    # 36 Name, 37 Type mand, 38 Save, 39 Back, 40 Delete, 41 Inserted By
    sit_reason = [
        [5, "SIT-006", "SKN-034", "Master Reason", "Type setelah Create",
         "View Type setelah create",
         "Value Type reason menjadi null padahal di DB ada datanya",
         "1. Add Reason, isi Name+Type, Save.\n2. Buka detail / list, cek Type.",
         "Gambar 8", "High", "In Progress", "Farrel (Maven)", "September 4, 2026", ""],
        [6, "SIT-007", "SKN-035", "Master Reason", "Type Visit / Others",
         "Add Type Visit dan Others",
         "Muncul error tidak enable input data dengan Tipe Reason Visit/Others",
         "1. Add Reason. Type Visit atau Others. Save.",
         "Gambar 9", "High", "Resolved", "Riyan (BE)", "September 4, 2026", ""],
        [7, "SIT-008", "SKN-032", "Master Reason", "Type Badge Warna dan Label",
         "Perbedaan warna UI pada setiap Type Reason",
         "Bedakan warna tiap Type Reason agar tidak rancu",
         "1. Buat reason beberapa type. Amati badge di tabel.",
         "Gambar 10", "Low", "Resolved", "Farrel (Maven)", "September 4, 2026", ""],
        [8, "SIT-009", "SKN-041", "Master Reason", "Inserted / Updated By",
         "Inserted by dan Updated by pada table Reason DB",
         "Sistem belum mencatat user pada Inserted By / Updated By di mGtReason",
         "1. Create/Update Reason. Cek kolom DB.",
         "Gambar 11", "High", "Open", "Riyan (BE)", "September 4, 2026", ""],
        [9, "SIT-010", "SKN-030", "Master Reason", "Filter Stat Cards - Type",
         "Validasi filter kolom Type (kartu Visit/Return/Order)",
         "Fungsi filter Type di dashboard Reason belum sesuai",
         "1. Klik kartu Visit. Pastikan hanya Kunjungan. Ulangi Return/Order.",
         "Gambar 12", "Medium", "In Progress", "Farrel (Maven)", "September 4, 2026", ""],
        [10, "SIT-011", "SKN-029", "Master Reason", "Filter Stat Cards - Status",
         "Validasi filter Status Active/Inactive",
         "Fungsi filter Status di dashboard Reason belum sesuai",
         "1. Filter Active lalu Inactive.",
         "Gambar 13", "Medium", "In Progress", "Farrel (Maven)", "September 4, 2026", ""],
        [11, "SIT-012", "SKN-030", "Master Reason", "Filter Stat Cards - Type",
         "Value Type Reason UI Maven vs backend Falcon FPRS",
         "Label Visit di UI vs nilai Kunjungan di backend",
         "1. Amati label Visit di UI. 2. Cek mGtReason.txtTipe = Kunjungan.",
         "Gambar 14", "High", "Open", "Farrel (Maven)", "September 4, 2026", ""],
    ]

    sit = wb.create_sheet("Issue SIT")
    sit["A1"] = "LOG ISSUE / DEFECT SYSTEM INTEGRATION TESTING (SIT)"
    sit["A1"].font = TITLE_FONT
    sit.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(SIT_HEADERS))
    sit["A2"] = (
        "Issue Product + Reason dari log SIT lama. Ref ID mengarah ke Scenario ID sheet Test Case. "
        "Tambah baris baru bila SKN Fail."
    )
    sit["A2"].alignment = WRAP
    sit.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(SIT_HEADERS))
    sit.row_dimensions[2].height = 32
    for c, h in enumerate(SIT_HEADERS, 1):
        sit.cell(3, c, h)
    style_header(sit, 3, len(SIT_HEADERS))
    sit_all = list(PRODUCT_SIT) + sit_reason
    write_rows(sit, 4, sit_all)
    for r in range(4, 4 + len(sit_all)):
        sit.row_dimensions[r].height = 68
    sit.freeze_panes = "A4"
    sit.auto_filter.ref = f"A3:N{3 + len(sit_all)}"
    autosize(sit, [6, 12, 12, 16, 32, 40, 42, 48, 12, 16, 14, 16, 20, 16])

    counts = {}
    for row in rows:
        counts[row[2]] = counts.get(row[2], 0) + 1
    summary = wb.create_sheet("Summary", 0)
    summary["A1"] = "MAN POWER GT — TEST CASE SUMMARY"
    summary["A1"].font = TITLE_FONT
    summary["A3"] = "Modul / Area"
    summary["B3"] = "Jumlah SKN"
    style_header(summary, 3, 2)
    for i, (mod, n) in enumerate(counts.items(), 4):
        summary.cell(i, 1, mod).font = CELL_FONT
        summary.cell(i, 2, n).font = CELL_FONT
    summary.cell(4 + len(counts), 1, "Total").font = Font(bold=True, name="Calibri")
    summary.cell(4 + len(counts), 2, len(rows)).font = Font(bold=True, name="Calibri")
    summary["A3"].fill = GREEN
    summary["B3"].fill = GREEN
    autosize(summary, [28, 16])

    wb.save(OUT)
    print(f"Wrote {OUT}: {len(rows)} SKN, {len(sit_all)} SIT")
    for mod, n in counts.items():
        print(f"  {mod}: {n}")


if __name__ == "__main__":
    main()
