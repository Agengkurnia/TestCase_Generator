"""Sample workbook: Master Product saja."""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
REPO = HERE.parent.parent
sys.path.insert(0, str(REPO))
sys.path.insert(0, str(HERE))

from cases_product import CASES, MOD  # noqa: E402
from xlsx_kit import (  # noqa: E402
    CELL_FONT,
    GREEN,
    SKN_HEADERS,
    SIT_HEADERS,
    TITLE_FONT,
    WRAP,
    autosize,
    style_header,
    write_rows,
)

OUT = HERE / "ManPowerGT_MasterProduct_TestCase_Sample.xlsx"


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Test Case (SKN)"
    ws["A1"] = "TEST CASE — MAN POWER GT — MASTER PRODUCT (SAMPLE)"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(SKN_HEADERS))
    ws["A2"] = "Sample Product. Workbook penuh: generate_masters.py / run_masters_sit.py."
    ws["A2"].alignment = WRAP
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(SKN_HEADERS))
    for col, header in enumerate(SKN_HEADERS, 1):
        ws.cell(3, col, header)
    style_header(ws, 3, len(SKN_HEADERS))
    rows = []
    for i, (field, tipe, desc, pre, steps, exp, pri) in enumerate(CASES, 1):
        rows.append(
            [i, f"SKN-{i:03d}", MOD, field, tipe, desc, pre, steps, exp, pri, "Not Run", "", "", ""]
        )
    write_rows(ws, 4, rows)
    for row in range(4, 4 + len(rows)):
        ws.row_dimensions[row].height = 78
    ws.auto_filter.ref = f"A3:N{3 + len(rows)}"
    ws.freeze_panes = "A4"
    autosize(ws, [6, 14, 18, 36, 12, 44, 40, 52, 52, 12, 12, 22, 14, 16])
    dv = DataValidation(type="list", formula1='"Not Run,Pass,Fail"', allow_blank=True)
    dv.add(f"K4:K{3 + len(rows)}")
    ws.add_data_validation(dv)

    sit = wb.create_sheet("Issue SIT")
    sit["A1"] = "LOG ISSUE / DEFECT SYSTEM INTEGRATION TESTING (SIT)"
    sit["A1"].font = TITLE_FONT
    sit.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(SIT_HEADERS))
    sit["A2"] = "Kosong. Sample Product tidak diisi runner penuh."
    sit["A2"].alignment = WRAP
    for col, header in enumerate(SIT_HEADERS, 1):
        sit.cell(3, col, header)
    style_header(sit, 3, len(SIT_HEADERS))
    sit.freeze_panes = "A4"
    autosize(sit, [6, 12, 12, 16, 32, 40, 42, 48, 12, 16, 14, 16, 20, 16])

    summary = wb.create_sheet("Summary", 0)
    summary["A1"] = "MAN POWER GT — PRODUCT SAMPLE"
    summary["A1"].font = TITLE_FONT
    summary["A3"] = "Modul / Area"
    summary["B3"] = "Jumlah SKN"
    style_header(summary, 3, 2)
    summary["A3"].fill = GREEN
    summary["B3"].fill = GREEN
    summary.cell(4, 1, MOD).font = CELL_FONT
    summary.cell(4, 2, len(rows)).font = CELL_FONT
    wb.save(OUT)
    print(f"Wrote {OUT.name}: {len(rows)} SKN")


if __name__ == "__main__":
    main()
