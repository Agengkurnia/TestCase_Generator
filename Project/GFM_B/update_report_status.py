"""Isi kategori outstanding vs pass di GFM_B_Testcase.xlsx."""
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parent.parent
sys.path.insert(0, str(REPO))
sys.path.insert(0, str(HERE))

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

from xlsx_kit import CELL_FONT, GREEN, TITLE_FONT, WRAP, autosize, style_header

FIXED = {
    "SKN-007": "Pass: filter kolom Inventory sudah ada",
    "SKN-024": "Pass: peta koordinat di detail Inventory sudah ada",
    "SKN-035": "Pass: mPDF terpasang, Print Report Inventory jalan",
    "SKN-088": "Pass: upload foto inspeksi/maintenance sudah ada",
}

BLOCKED = {
    "SKN-005",
    "SKN-030",
    "SKN-046",
    "SKN-051",
    "SKN-136",
}


def main() -> None:
    wb = load_workbook(HERE / "GFM_B_Testcase.xlsx")
    ws = wb["Test Case (SKN)"]
    ws.cell(3, 16, "Kategori")
    style_header(ws, 3, 16)
    ws.column_dimensions["P"].width = 28

    passed = failed = blocked = 0
    outstanding = []
    for r in range(4, ws.max_row + 1):
        sid = ws.cell(r, 2).value
        status = ws.cell(r, 11).value
        if sid in BLOCKED or status == "Not Run":
            cat = "Blocked (bukan defect)"
            blocked += 1
        elif status == "Fail":
            cat = "Outstanding"
            failed += 1
            outstanding.append(
                (
                    sid,
                    ws.cell(r, 3).value,
                    ws.cell(r, 4).value,
                    str(ws.cell(r, 12).value or ""),
                )
            )
        else:
            cat = "Pass"
            passed += 1
            if sid in FIXED:
                ws.cell(r, 12, FIXED[sid])
                ws.cell(r, 12).alignment = Alignment(wrap_text=True, vertical="top")
        cell = ws.cell(r, 16, cat)
        cell.font = CELL_FONT
        cell.alignment = Alignment(vertical="top")

    summary = wb["Summary"]
    for r in range(23, 90):
        for c in range(1, 5):
            summary.cell(r, c, None)

    summary["A23"] = "STATUS SETELAH PERBAIKAN GAP PRODUK"
    summary["A23"].font = TITLE_FONT
    summary.merge_cells("A23:D23")

    summary["A24"] = "Kategori"
    summary["B24"] = "Jumlah"
    summary["C24"] = "Arti"
    style_header(summary, 24, 3)
    rows = [
        ("Pass", passed, "SKN lulus, termasuk perbaikan filter/peta/PDF/foto (007, 024, 035, 088)."),
        ("Blocked (bukan defect)", blocked, "Tidak diuji: dataset SIT tidak kosong, atau uji 21MB manual."),
        ("Outstanding", failed, "Masih Fail. Bukan gap yang sudah diperbaiki; runner/data atau kasus lain."),
    ]
    for i, (label, n, note) in enumerate(rows, 25):
        summary.cell(i, 1, label).font = CELL_FONT
        summary.cell(i, 2, n).font = CELL_FONT
        summary.cell(i, 3, note).font = CELL_FONT
        summary.cell(i, 3).alignment = WRAP

    summary["A29"] = "Sudah Pass (perbaikan 13 Sep 2026)"
    summary["A29"].font = Font(bold=True, name="Calibri")
    summary["A30"] = "SKN"
    summary["B30"] = "Keterangan"
    style_header(summary, 30, 2)
    for i, (sid, note) in enumerate(FIXED.items(), 31):
        summary.cell(i, 1, sid).font = CELL_FONT
        summary.cell(i, 2, note).font = CELL_FONT

    summary["A36"] = "Outstanding Fail"
    summary["A36"].font = Font(bold=True, name="Calibri")
    summary["A37"] = "SKN"
    summary["B37"] = "Modul"
    summary["C37"] = "Field"
    summary["D37"] = "Hasil aktual"
    style_header(summary, 37, 4)
    if not outstanding:
        summary.cell(38, 1, "(tidak ada)").font = CELL_FONT
    else:
        for i, (sid, mod, field, actual) in enumerate(outstanding, 38):
            summary.cell(i, 1, sid).font = CELL_FONT
            summary.cell(i, 2, mod).font = CELL_FONT
            summary.cell(i, 3, field).font = CELL_FONT
            summary.cell(i, 4, actual).font = CELL_FONT
            summary.cell(i, 4).alignment = WRAP
            summary.row_dimensions[i].height = 36

    summary["A22"] = (
        "Fail yang tersisa BUKAN SKN-007/024/035/088. "
        "Itu kasus lain (wizard inspeksi, Priority Alpine, Preservation, Mitigation, Map, CMS) "
        "atau Blocked karena data SIT tidak kosong."
    )
    summary["A22"].alignment = WRAP
    summary.merge_cells("A22:D22")
    summary.row_dimensions[22].height = 40
    autosize(summary, [28, 22, 36, 72])
    summary["A24"].fill = GREEN
    summary["B24"].fill = GREEN
    summary["C24"].fill = GREEN

    wb.save(HERE / "GFM_B_Testcase.xlsx")
    print(f"Updated {HERE / 'GFM_B_Testcase.xlsx'}  Pass={passed} Blocked={blocked} Outstanding={failed}")
    for sid, mod, field, _ in outstanding:
        print(f"  {sid} {mod} {field}")


if __name__ == "__main__":
    main()
