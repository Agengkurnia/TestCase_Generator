"""Workbook test case: GFM_B — semua modul operasional + CMS + Access Control."""
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parent.parent
sys.path.insert(0, str(REPO))
sys.path.insert(0, str(HERE))

from cases_access import CASES as ACCESS_CASES  # noqa: E402
from cases_access import MOD as MOD_ACCESS  # noqa: E402
from cases_cms import CASES as CMS_CASES  # noqa: E402
from cases_cms import MOD as MOD_CMS  # noqa: E402
from cases_dashboard import CASES as DASHBOARD_CASES  # noqa: E402
from cases_dashboard import MOD as MOD_DASHBOARD  # noqa: E402
from cases_inspection import CASES as INSPECTION_CASES  # noqa: E402
from cases_inspection import MOD as MOD_INSPECTION  # noqa: E402
from cases_map import CASES as MAP_CASES  # noqa: E402
from cases_map import MOD as MOD_MAP  # noqa: E402
from cases_mitigation import CASES as MITIGATION_CASES  # noqa: E402
from cases_mitigation import MOD as MOD_MITIGATION  # noqa: E402
from cases_preservation import CASES as PRESERVATION_CASES  # noqa: E402
from cases_preservation import MOD as MOD_PRESERVATION  # noqa: E402
from cases_priority import CASES as PRIORITY_CASES  # noqa: E402
from cases_priority import MOD as MOD_PRIORITY  # noqa: E402
from generate_inventory import CASES as INVENTORY_CASES  # noqa: E402
from generate_inventory import MOD as MOD_INVENTORY  # noqa: E402
from xlsx_kit import (  # noqa: E402
    CELL_FONT,
    GREEN,
    SKN_HEADERS,
    SIT_HEADERS,
    TITLE_FONT,
    WRAP,
    autosize,
    style_header,
    write_rows,
)
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

OUT = HERE / "GFM_B_Testcase.xlsx"

MODULES = [
    (MOD_INVENTORY, INVENTORY_CASES),
    (MOD_DASHBOARD, DASHBOARD_CASES),
    (MOD_PRIORITY, PRIORITY_CASES),
    (MOD_INSPECTION, INSPECTION_CASES),
    (MOD_PRESERVATION, PRESERVATION_CASES),
    (MOD_MITIGATION, MITIGATION_CASES),
    (MOD_MAP, MAP_CASES),
    (MOD_CMS, CMS_CASES),
    (MOD_ACCESS, ACCESS_CASES),
]


def packed():
    rows = []
    for mod, cases in MODULES:
        for field, tipe, desc, pre, steps, exp, pri in cases:
            rows.append((mod, field, tipe, desc, pre, steps, exp, pri))
    return rows


def main():
    skn = packed()
    wb = Workbook()

    ws = wb.active
    ws.title = "Test Case (SKN)"
    ws["A1"] = "TEST CASE — DASHBOARD SLOPE GFM"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(SKN_HEADERS))
    ws["A2"] = (
        "Sumber: docs/brd/01-dashboard.md s/d 09-access-control.md plus controller/view terkait. "
        "Filter kolom Modul / Area untuk satu modul. Inventory SKN-001 dst. tetap di awal "
        "agar runner SIT Inventory tidak bergeser. "
        "Kolom Status Uji / Hasil Aktual / QA / Tanggal Uji diisi saat eksekusi. "
        "Sheet Issue SIT: isi baris baru bila SKN Fail; Ref ID = Scenario ID di sheet ini."
    )
    ws["A2"].alignment = WRAP
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(SKN_HEADERS))
    ws.row_dimensions[2].height = 56

    for col, header in enumerate(SKN_HEADERS, 1):
        ws.cell(3, col, header)
    style_header(ws, 3, len(SKN_HEADERS))

    skn_rows = []
    for i, (mod, field, tipe, desc, pre, steps, exp, pri) in enumerate(skn, 1):
        skn_rows.append(
            [
                i,
                f"SKN-{i:03d}",
                mod,
                field,
                tipe,
                desc,
                pre,
                steps,
                exp,
                pri,
                "Not Run",
                "",
                "",
                "",
            ]
        )
    write_rows(ws, 4, skn_rows)
    for row in range(4, 4 + len(skn_rows)):
        ws.row_dimensions[row].height = 78
    ws.auto_filter.ref = f"A3:N{3 + len(skn_rows)}"
    ws.freeze_panes = "A4"
    ws.row_dimensions[3].height = 22
    autosize(ws, [6, 14, 16, 36, 12, 44, 40, 52, 52, 12, 12, 22, 14, 16])

    dv_status = DataValidation(type="list", formula1='"Not Run,Pass,Fail"', allow_blank=True)
    dv_status.add(f"K4:K{3 + len(skn_rows)}")
    ws.add_data_validation(dv_status)
    dv_tipe = DataValidation(type="list", formula1='"Positive,Negative"', allow_blank=False)
    dv_tipe.add(f"E4:E{3 + len(skn_rows)}")
    ws.add_data_validation(dv_tipe)

    sit = wb.create_sheet("Issue SIT")
    sit["A1"] = "LOG ISSUE / DEFECT SYSTEM INTEGRATION TESTING (SIT)"
    sit["A1"].font = TITLE_FONT
    sit.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(SIT_HEADERS))
    sit["A2"] = (
        "Belum ada log SIT. Tambah baris saat SKN berstatus Fail. "
        "Ref ID mengarah ke Scenario ID (SKN-xxx) di sheet Test Case."
    )
    sit["A2"].alignment = WRAP
    sit.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(SIT_HEADERS))
    sit.row_dimensions[2].height = 32
    for col, header in enumerate(SIT_HEADERS, 1):
        sit.cell(3, col, header)
    style_header(sit, 3, len(SIT_HEADERS))
    sit.freeze_panes = "A4"
    sit.auto_filter.ref = "A3:N3"
    autosize(sit, [6, 12, 12, 16, 32, 40, 42, 48, 12, 16, 14, 16, 20, 16])

    counts = {}
    for row in skn_rows:
        counts[row[2]] = counts.get(row[2], 0) + 1
    pos = sum(1 for row in skn_rows if row[4] == "Positive")
    neg = sum(1 for row in skn_rows if row[4] == "Negative")
    high = sum(1 for row in skn_rows if row[9] == "High")
    med = sum(1 for row in skn_rows if row[9] == "Medium")
    low = sum(1 for row in skn_rows if row[9] == "Low")

    summary = wb.create_sheet("Summary", 0)
    summary["A1"] = "GFM_B — TEST CASE SUMMARY"
    summary["A1"].font = TITLE_FONT
    summary["A3"] = "Modul / Area"
    summary["B3"] = "Jumlah SKN"
    style_header(summary, 3, 2)
    r = 4
    for mod, n in counts.items():
        summary.cell(r, 1, mod).font = CELL_FONT
        summary.cell(r, 2, n).font = CELL_FONT
        r += 1
    extras = [
        ("Positive", pos),
        ("Negative", neg),
        ("Priority High", high),
        ("Priority Medium", med),
        ("Priority Low", low),
        ("Issue SIT (awal)", 0),
    ]
    r += 1
    summary.cell(r, 1, "Tipe / Priority").font = Font(bold=True, name="Calibri")
    r += 1
    for label, n in extras:
        summary.cell(r, 1, label).font = CELL_FONT
        summary.cell(r, 2, n).font = CELL_FONT
        r += 1
    summary.cell(r, 1, "Total SKN").font = Font(bold=True, name="Calibri")
    summary.cell(r, 2, len(skn_rows)).font = Font(bold=True, name="Calibri")
    summary["A3"].fill = GREEN
    summary["B3"].fill = GREEN
    autosize(summary, [32, 16])

    wb.save(OUT)
    print(f"Wrote {OUT.name}: {len(skn_rows)} SKN")
    for mod, n in counts.items():
        print(f"  {mod}: {n}")
    print(f"  Positive={pos} Negative={neg} High={high} Medium={med} Low={low}")


if __name__ == "__main__":
    main()
