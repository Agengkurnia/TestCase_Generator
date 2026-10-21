"""Uji Print Report di produksi. Hanya login dan GET. Tidak membuat atau menghapus data.

Jalankan nanti dari folder ini:

    py -3 run_pdf_prod.py

Akun dibaca dari config.prod.json. Hasil: prod_pdf/GFM_B_PDF_Prod.xlsx dan prod_pdf/evidence/.
"""
from __future__ import annotations

import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))

from evidence_report import compose_card  # noqa: E402

CONFIG = json.loads((ROOT / "config.prod.json").read_text(encoding="utf-8"))
BASE = CONFIG["base_url"].rstrip("/")
EMAIL = CONFIG["email"]
PASSWORD = CONFIG["password"]
TESTER = CONFIG.get("tester", "Ageng")
TODAY = date.today().strftime("%d/%m/%Y")

OUT = ROOT / "prod_pdf"
RAW = OUT / "raw"
EVIDENCE = OUT / "evidence"
XLSX = OUT / "GFM_B_PDF_Prod.xlsx"

LINK_FONT = Font(name="Calibri", size=10, color="0563C1", underline="single")
CELL_FONT = Font(name="Calibri", size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
HEAD = Font(name="Calibri", size=11, bold=True, color="FFFFFF")


def highlight_target(page, selector: str, text_label: str) -> None:
    page.evaluate(
        """([sel, label]) => {
            document.querySelectorAll('.qa-test-highlight').forEach(el => el.remove());
            document.querySelectorAll('[data-qa-highlighted]').forEach(el => {
                el.style.outline = '';
                el.style.outlineOffset = '';
                el.style.boxShadow = '';
                el.removeAttribute('data-qa-highlighted');
            });
            const el = document.querySelector(sel);
            if (!el) return;
            el.setAttribute('data-qa-highlighted', 'true');
            el.scrollIntoView({ block: 'center', inline: 'center' });
            window.scrollBy(0, -80);
            el.style.outline = '4px solid #E53935';
            el.style.outlineOffset = '2px';
            el.style.boxShadow = '0 0 18px rgba(229, 57, 53, 0.95)';
            const rect = el.getBoundingClientRect();
            const badge = document.createElement('div');
            badge.className = 'qa-test-highlight';
            badge.innerText = '🎯 TEST TARGET: ' + label;
            badge.style.position = 'fixed';
            badge.style.top = Math.max(10, rect.top - 30) + 'px';
            badge.style.left = Math.max(10, rect.left) + 'px';
            badge.style.background = '#E53935';
            badge.style.color = '#FFFFFF';
            badge.style.padding = '4px 10px';
            badge.style.borderRadius = '4px';
            badge.style.fontSize = '12px';
            badge.style.fontWeight = 'bold';
            badge.style.zIndex = '999999';
            badge.style.boxShadow = '0 2px 8px rgba(0,0,0,0.5)';
            badge.style.fontFamily = 'Calibri, sans-serif';
            document.body.appendChild(badge);
        }""",
        [selector, text_label],
    )
    page.wait_for_timeout(400)


def shot(page, skn: int) -> Path:
    RAW.mkdir(parents=True, exist_ok=True)
    path = RAW / f"SKN-{skn:03d}.png"
    page.screenshot(path=str(path), full_page=False)
    return path


def login(page) -> None:
    page.context.clear_cookies()
    page.goto(f"{BASE}/login", wait_until="domcontentloaded")
    page.fill("#email", EMAIL)
    page.fill("#password", PASSWORD)
    page.click("button[type=submit]")
    page.wait_for_load_state("domcontentloaded")
    if "/choose-role" in page.url:
        button = page.locator("button[type=submit], button:has-text('Continue')").first
        button.click()
        page.wait_for_load_state("domcontentloaded")
    if "/login" in page.url:
        raise SystemExit("Login produksi gagal. Cek email, password, atau halaman login.")


def pdf_ok(response) -> tuple[bool, str]:
    ctype = response.headers.get("content-type", "")
    body = response.body()
    magic = body[:5] == b"%PDF-"
    if response.status == 200 and "pdf" in ctype and magic and "/login" not in response.url:
        return True, f"HTTP 200, {ctype}, ukuran {len(body)} byte, header %PDF. URL {response.url}"
    return False, f"HTTP {response.status}, content-type {ctype or '-'}, magic PDF={magic}, URL {response.url}"


def inventory_targets(page) -> list[tuple[str, str]]:
    page.goto(f"{BASE}/inventory", wait_until="domcontentloaded")
    rows = page.locator("tbody tr")
    picked: dict[str, str] = {}
    for i in range(rows.count()):
        row = rows.nth(i)
        link = row.locator("a:has-text('Details')")
        if link.count() == 0:
            continue
        href = link.first.get_attribute("href") or ""
        if "/inventory/" not in href:
            continue
        kind = row.inner_text().lower()
        key = "rock" if "rock" in kind else "cut" if "cut" in kind else "other"
        picked.setdefault(key, href)
        if "rock" in picked and ("cut" in picked or "other" in picked):
            break
    return list(picked.items())


def first_href(page, selector: str) -> str:
    link = page.locator(selector).first
    if link.count() == 0:
        return ""
    return link.get_attribute("href") or ""


def check_print(page, skn: int, url: str, selector: str, label: str) -> dict:
    page.goto(url, wait_until="domcontentloaded")
    link = page.locator(selector).first
    if link.count() == 0:
        path = shot(page, skn)
        return {
            "status": "Not Run",
            "actual": f"Tombol Print Report tidak ada di {url}. Skrip tidak membuat data.",
            "shot": path,
            "points": ["Response PDF (Content-Type application/pdf).", "Isi file diawali %PDF."],
        }
    href = link.get_attribute("href") or ""
    highlight_target(page, selector, label)
    path = shot(page, skn)
    response = page.request.get(href)
    ok, note = pdf_ok(response)
    return {
        "status": "Pass" if ok else "Fail",
        "actual": note,
        "shot": path,
        "points": ["Response PDF (Content-Type application/pdf).", "Isi file diawali %PDF."],
    }


def run(page) -> list[dict]:
    login(page)
    cases = []

    targets = inventory_targets(page)
    if not targets:
        page.goto(f"{BASE}/inventory", wait_until="domcontentloaded")
        cases.append({
            "skn": 1,
            "module": "Inventory",
            "field": "PDF - Print Report",
            "status": "Fail",
            "actual": "Daftar Inventory kosong atau tautan Details tidak ditemukan.",
            "shot": shot(page, 1),
            "points": ["Ada slope yang bisa dibuka.", "Print Report mengembalikan PDF."],
        })
    else:
        notes = []
        ok_all = True
        page.goto(targets[0][1], wait_until="domcontentloaded")
        highlight_target(page, "a[href*='pdf-inventory']", "Print Report Inventory")
        shot_path = shot(page, 1)
        for kind, href in targets:
            page.goto(href, wait_until="domcontentloaded")
            pdf_href = first_href(page, "a[href*='pdf-inventory']")
            if not pdf_href:
                ok_all = False
                notes.append(f"{kind}: Print Report tidak ada.")
                continue
            ok, note = pdf_ok(page.request.get(pdf_href))
            ok_all = ok_all and ok
            notes.append(f"{kind}: {note}")
        cases.append({
            "skn": 1,
            "module": "Inventory",
            "field": "PDF - Print Report",
            "status": "Pass" if ok_all else "Fail",
            "actual": " ".join(notes),
            "shot": shot_path,
            "points": ["Print Report terbuka dari detail slope.", "Response application/pdf dan diawali %PDF."],
        })

    page.goto(f"{BASE}/priority", wait_until="domcontentloaded")
    detail = first_href(page, "a[href*='/priority/']")
    if not detail:
        cases.append({
            "skn": 2,
            "module": "Priority",
            "field": "PDF - Print Report",
            "status": "Not Run",
            "actual": "Tidak ada slope di Priority. Skrip tidak membuat data.",
            "shot": shot(page, 2),
            "points": ["Detail Priority punya Print Report.", "Response application/pdf dan diawali %PDF."],
        })
    else:
        item = check_print(page, 2, detail, "a[href*='pdf-rating']", "Print Report Rating")
        item.update({"skn": 2, "module": "Priority", "field": "PDF - Print Report"})
        cases.append(item)

    page.goto(f"{BASE}/mitigation", wait_until="domcontentloaded")
    calculators = page.locator("a:has-text('Open Calculator')")
    found = ""
    limit = min(calculators.count(), 8)
    for i in range(limit):
        href = calculators.nth(i).get_attribute("href") or ""
        if not href:
            continue
        page.goto(href, wait_until="domcontentloaded")
        report = first_href(page, "a:has-text('Print Report')")
        if report:
            found = report if report.startswith("http") else BASE + report if report.startswith("/") else page.url
            break
    if not found:
        cases.append({
            "skn": 3,
            "module": "Mitigation",
            "field": "PDF - Print Report",
            "status": "Not Run",
            "actual": "Tidak ada plan aktif pada 8 slope pertama. Skrip tidak membuat plan.",
            "shot": shot(page, 3),
            "points": ["Plan aktif punya Print Report.", "Response application/pdf dan diawali %PDF."],
        })
    else:
        item = check_print(page, 3, found, "a[href*='pdf-mitigation']", "Print Report Mitigation")
        item.update({"skn": 3, "module": "Mitigation", "field": "PDF - Print Report"})
        cases.append(item)
    return cases


def write_rows(cases: list[dict]) -> None:
    EVIDENCE.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "PDF Prod"
    headers = [
        "No.", "Scenario ID", "Modul / Area", "Field Skenario", "Tipe", "Deskripsi",
        "Prasyarat", "Langkah", "Expected", "Priority", "Status Uji", "Hasil Aktual",
        "Bukti Pengujian", "QA", "Tanggal Uji",
    ]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col, header)
        cell.font = HEAD
        cell.fill = PatternFill("solid", fgColor="1E7A91")
    for case in cases:
        skn = case["skn"]
        points = case["points"]
        mark = {"Pass": "Sesuai", "Fail": "Tidak terbukti"}.get(case["status"], "Tidak diuji")
        lines = [f"{i}. [{mark}] {point}" for i, point in enumerate(points, 1)]
        lines += ["", f"Catatan: {case['actual']}"]
        actual = "\n".join(lines)
        card = EVIDENCE / f"SKN-{skn:03d}.png"
        compose_card(
            card,
            skn=skn,
            status=case["status"],
            module=case["module"],
            field=case["field"],
            points=points,
            note=case["actual"],
            raw=case["shot"],
        )
        row = skn + 1
        values = [
            skn, f"PROD-PDF-{skn:03d}", case["module"], case["field"], "SIT",
            "Unduh PDF laporan di produksi tanpa mengubah data.",
            f"Login {EMAIL} di {BASE}.",
            "Buka modul, tandai Print Report, unduh PDF.",
            "\n".join(points),
            "High",
            case["status"],
            actual,
            f"evidence/SKN-{skn:03d}.png",
            TESTER,
            TODAY,
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row, col, value)
            cell.font = CELL_FONT
            cell.alignment = WRAP
        evid = ws.cell(row, 13)
        evid.hyperlink = f"evidence/SKN-{skn:03d}.png"
        evid.font = LINK_FONT
        ws.row_dimensions[row].height = 72
        print(f"PROD-PDF-{skn:03d} {case['status']} {case['actual']}")
    ws.column_dimensions["L"].width = 72
    ws.auto_filter.ref = f"A1:O{len(cases) + 1}"
    wb.save(XLSX)
    print(f"Excel: {XLSX}")


def main() -> None:
    print(f"Prod PDF check {BASE}. Hanya login dan unduh. Tidak membuat atau menghapus data.")
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        context = browser.new_context(viewport={"width": 1400, "height": 900})
        page = context.new_page()
        page.set_default_timeout(45000)
        try:
            cases = run(page)
        finally:
            context.close()
            browser.close()
    write_rows(cases)
    failed = [c["skn"] for c in cases if c["status"] == "Fail"]
    if failed:
        raise SystemExit(f"Fail: {failed}")


if __name__ == "__main__":
    main()
