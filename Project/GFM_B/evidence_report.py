"""Kartu evidence per SKN: poin expected + screenshot, tanpa thumbnail numpuk di Excel."""
from __future__ import annotations

import re
import shutil
from pathlib import Path

from openpyxl.styles import Alignment, Font
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
EVIDENCE = ROOT / "evidence"
RAW = EVIDENCE / "raw"

CARD_W = 1100
PAD = 28
SHOT_MAX_H = 520
SIDEBAR_CROP = 256

STATUS_BAR = {
    "Pass": (30, 125, 55),
    "Fail": (180, 35, 24),
    "Not Run": (181, 71, 8),
}
STATUS_LABEL = {
    "Pass": "PASS",
    "Fail": "FAIL",
    "Not Run": "TIDAK DIUJI",
}
MARK = {
    "Pass": ("Sesuai", (30, 125, 55)),
    "Fail": ("Tidak terbukti", (180, 35, 24)),
    "Not Run": ("Tidak diuji", (181, 71, 8)),
}

LINK_FONT = Font(name="Calibri", size=10, color="0563C1", underline="single")
CELL_FONT = Font(name="Calibri", size=10)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")


def split_points(text: str | None) -> list[str]:
    if not text:
        return []
    out = []
    for raw in str(text).replace("\r\n", "\n").split("\n"):
        line = raw.strip()
        if not line:
            continue
        line = re.sub(r"^\d+[\.\)]\s*", "", line)
        out.append(line)
    return out


def extract_note(actual: str | None) -> str:
    text = (actual or "").strip()
    if not text:
        return ""
    if "Catatan:" in text:
        return text.split("Catatan:", 1)[1].strip()
    if re.match(r"^1\.\s*\[", text):
        return ""
    return text


def format_actual(status: str, points: list[str], note: str) -> str:
    mark, _ = MARK.get(status, MARK["Not Run"])
    if not points:
        return note or f"Tidak ada poin expected. Status: {status}."
    lines = [f"{i}. [{mark}] {p}" for i, p in enumerate(points, 1)]
    if note:
        lines.append("")
        lines.append(f"Catatan: {note}")
    return "\n".join(lines)


def _font(size: int, bold: bool = False):
    names = (
        ("segoeuib.ttf", "segoeui.ttf")
        if bold
        else ("segoeui.ttf", "calibri.ttf", "arial.ttf")
    )
    if bold:
        names = ("segoeuib.ttf", "calibrib.ttf", "arialbd.ttf", "segoeui.ttf")
    for name in names:
        path = Path(r"C:\Windows\Fonts") / name
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def _wrap(draw: ImageDraw.ImageDraw, text: str, font, max_w: int) -> list[str]:
    words = text.split()
    if not words:
        return [""]
    lines, cur = [], words[0]
    for word in words[1:]:
        trial = f"{cur} {word}"
        if draw.textlength(trial, font=font) <= max_w:
            cur = trial
        else:
            lines.append(cur)
            cur = word
    lines.append(cur)
    return lines


def _avg_rgb(im: Image.Image) -> tuple[int, int, int]:
    small = im.convert("RGB").resize((8, 8))
    r = g = b = n = 0
    for y in range(8):
        for x in range(8):
            pr, pg, pb = small.getpixel((x, y))
            r += pr
            g += pg
            b += pb
            n += 1
    return (r // n, g // n, b // n)


def crop_main(im: Image.Image) -> Image.Image:
    w, h = im.size
    if w < 900:
        return im
    avg = _avg_rgb(im.crop((0, 0, 40, h)))
    if avg[0] < 90 and avg[1] < 100 and avg[2] < 120:
        return im.crop((SIDEBAR_CROP, 0, w, h))
    return im


def ensure_raw(skn: int, preferred: Path | None = None) -> Path | None:
    RAW.mkdir(parents=True, exist_ok=True)
    dest = RAW / f"SKN-{skn:03d}.png"
    if preferred and preferred.exists() and preferred.resolve() != dest.resolve():
        shutil.copy2(preferred, dest)
        return dest
    if dest.exists():
        return dest
    src = EVIDENCE / f"SKN-{skn:03d}.png"
    if src.exists():
        shutil.copy2(src, dest)
        return dest
    return None


def compose_card(
    out: Path,
    *,
    skn: int,
    status: str,
    module: str,
    field: str,
    points: list[str],
    note: str,
    raw: Path | None,
) -> None:
    title_f = _font(22, bold=True)
    meta_f = _font(14, bold=True)
    body_f = _font(15)
    small_f = _font(13)
    scratch = Image.new("RGB", (CARD_W, 200), (255, 255, 255))
    draw = ImageDraw.Draw(scratch)
    inner = CARD_W - PAD * 2

    mark, _mark_rgb = MARK.get(status, MARK["Not Run"])
    bar = STATUS_BAR.get(status, STATUS_BAR["Not Run"])
    sid = f"SKN-{skn:03d}"
    blocks: list[tuple[str, object, tuple[int, int, int], int]] = []

    def add(text: str, font, color, indent: int = 0):
        if not text:
            blocks.append(("", font, color, 0))
            return
        for line in _wrap(draw, text, font, inner - indent):
            blocks.append((line, font, color, indent))

    add("Dicek terhadap Hasil yang Diharapkan", meta_f, (31, 41, 55))
    add("", body_f, (0, 0, 0))
    if points:
        for i, point in enumerate(points, 1):
            add(f"{i}. [{mark}] {point}", body_f, (17, 24, 39))
        add("", body_f, (0, 0, 0))
    else:
        add("(Tidak ada poin expected terpisah.)", body_f, (75, 85, 99))
        add("", body_f, (0, 0, 0))

    if note:
        add("Catatan uji", meta_f, (31, 41, 55))
        add(note, body_f, (55, 65, 81))
        add("", body_f, (0, 0, 0))

    add("Tampilan saat uji", meta_f, (31, 41, 55))

    shot = None
    if raw and raw.exists():
        shot = crop_main(Image.open(raw).convert("RGB"))
        scale = min(inner / shot.width, SHOT_MAX_H / shot.height, 1.0)
        shot = shot.resize(
            (max(1, int(shot.width * scale)), max(1, int(shot.height * scale))),
            Image.Resampling.LANCZOS,
        )

    line_h = 22
    header_h = 72
    shot_h = (shot.height + 16) if shot else 36
    text_h = sum(line_h if t else 10 for t, _, _, _ in blocks)
    height = PAD + header_h + text_h + shot_h + 36 + PAD

    card = Image.new("RGB", (CARD_W, height), (255, 255, 255))
    d = ImageDraw.Draw(card)
    d.rectangle((0, 0, CARD_W, header_h), fill=bar)
    d.text(
        (PAD, 14),
        f"{sid}    {STATUS_LABEL.get(status, status.upper())}",
        font=title_f,
        fill=(255, 255, 255),
    )
    d.text((PAD, 44), f"{module}  ·  {field}", font=meta_f, fill=(235, 245, 235))

    y = header_h + 16
    for text, font, color, indent in blocks:
        if not text:
            y += 10
            continue
        d.text((PAD + indent, y), text, font=font, fill=color)
        y += line_h

    y += 4
    if shot:
        card.paste(shot, (PAD, y))
        d.rectangle((PAD - 1, y - 1, PAD + shot.width, y + shot.height), outline=(209, 213, 219))
        y += shot.height + 10
    else:
        d.text((PAD, y), "Screenshot tidak tersedia untuk SKN ini.", font=small_f, fill=(107, 114, 128))
        y += 24

    raw_name = f"evidence/raw/SKN-{skn:03d}.png" if raw else "-"
    d.text((PAD, height - 28), f"File screenshot: {raw_name}", font=small_f, fill=(107, 114, 128))
    card.save(out, "PNG", optimize=True)


def strip_drawings(ws) -> None:
    ws._images.clear()


def refresh_evidence(ws, *, results: dict | None = None, total: int = 157) -> None:
    strip_drawings(ws)
    if ws.cell(3, 15).value != "Evidence":
        ws.cell(3, 15, "Evidence")
    ws.column_dimensions["O"].width = 22
    ws.column_dimensions["L"].width = 56

    for i in range(1, total + 1):
        row = 3 + i
        status = None
        note = ""
        preferred = None
        if results and i in results:
            status = results[i]["status"]
            note = results[i]["actual"]
            shot = results[i].get("shot")
            preferred = Path(shot) if shot else None
        else:
            status = ws.cell(row, 11).value
            note = extract_note(ws.cell(row, 12).value)
        if not status:
            continue
        status = str(status)
        points = split_points(ws.cell(row, 9).value)
        actual = format_actual(status, points, note)
        cell_l = ws.cell(row, 12, actual)
        cell_l.alignment = WRAP_TOP
        cell_l.font = CELL_FONT

        raw = ensure_raw(i, preferred)
        card = EVIDENCE / f"SKN-{i:03d}.png"
        compose_card(
            card,
            skn=i,
            status=status,
            module=str(ws.cell(row, 3).value or ""),
            field=str(ws.cell(row, 4).value or ""),
            points=points,
            note=note,
            raw=raw,
        )
        label = f"SKN-{i:03d}.png\n{len(points)} poin · {status}"
        evid = ws.cell(row, 15, label)
        evid.hyperlink = f"evidence/SKN-{i:03d}.png"
        evid.font = LINK_FONT
        evid.alignment = WRAP_TOP
        n_lines = actual.count("\n") + 2
        ws.row_dimensions[row].height = min(200, max(48, 14 * min(n_lines, 12) + 10))


def rewrite_workbook(xlsx: Path, *, results: dict | None = None, total: int = 157) -> Path:
    from openpyxl import load_workbook

    EVIDENCE.mkdir(exist_ok=True)
    wb = load_workbook(xlsx)
    refresh_evidence(wb["Test Case (SKN)"], results=results, total=total)
    if "Issue SIT" in wb.sheetnames:
        strip_drawings(wb["Issue SIT"])
    saved = xlsx
    try:
        wb.save(xlsx)
    except PermissionError:
        saved = xlsx.with_name(xlsx.stem + "_Hasil.xlsx")
        wb.save(saved)
    for leftover in EVIDENCE.glob("*_xl.png"):
        leftover.unlink(missing_ok=True)
    return saved


if __name__ == "__main__":
    out = rewrite_workbook(ROOT / "GFM_B_Testcase.xlsx")
    print(f"Evidence rewritten: {out}")
