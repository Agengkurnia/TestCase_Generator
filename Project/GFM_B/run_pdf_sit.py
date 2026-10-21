"""SIT Print Report: Inventory, Priority, Mitigation. Kotak merah + kartu evidence."""
from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(REPO))

from evidence_report import EVIDENCE, RAW, compose_card, ensure_raw, split_points  # noqa: E402
from run_inventory_sit import BASE, TESTER, login, shot  # noqa: E402

XLSX = ROOT / "GFM_B_Testcase.xlsx"
TODAY = date.today().strftime("%d/%m/%Y")
VIEWER = ("sit.viewer@gfm.test", "password")
ADMIN = ("sit.admin@gfm.test", "password")
LINK_FONT = Font(name="Calibri", size=10, color="0563C1", underline="single")
CELL_FONT = Font(name="Calibri", size=10)
WRAP = Alignment(wrap_text=True, vertical="top")


def highlight_target(page, selector: str, text_label: str) -> None:
    page.evaluate(
        """([sel, label]) => {
            document.querySelectorAll('.qa-test-highlight').forEach(el => el.remove());
            document.querySelectorAll('[data-qa-highlighted]').forEach(el => {
                el.style.outline = '';
                el.style.outlineOffset = '';
                el.style.boxShadow = '';
                el.removeAttribute('data-qa-highlighted');
            });
            const el = document.querySelector(sel);
            if (!el) return;
            el.setAttribute('data-qa-highlighted', 'true');
            el.scrollIntoView({ block: 'center', inline: 'center' });
            window.scrollBy(0, -80);
            el.style.outline = '4px solid #E53935';
            el.style.outlineOffset = '2px';
            el.style.boxShadow = '0 0 18px rgba(229, 57, 53, 0.95)';
            const rect = el.getBoundingClientRect();
            const badge = document.createElement('div');
            badge.className = 'qa-test-highlight';
            badge.innerText = '🎯 TEST TARGET: ' + label;
            badge.style.position = 'fixed';
            badge.style.top = Math.max(10, rect.top - 30) + 'px';
            badge.style.left = Math.max(10, rect.left) + 'px';
            badge.style.background = '#E53935';
            badge.style.color = '#FFFFFF';
            badge.style.padding = '4px 10px';
            badge.style.borderRadius = '4px';
            badge.style.fontSize = '12px';
            badge.style.fontWeight = 'bold';
            badge.style.zIndex = '999999';
            badge.style.boxShadow = '0 2px 8px rgba(0,0,0,0.5)';
            badge.style.fontFamily = 'Calibri, sans-serif';
            document.body.appendChild(badge);
        }""",
        [selector, text_label],
    )
    page.wait_for_timeout(400)


def pdf_ok(response) -> tuple[bool, str]:
    ctype = response.headers.get("content-type", "")
    body = response.body()
    magic = body[:5] == b"%PDF-"
    if response.status == 200 and "pdf" in ctype and magic:
        return True, f"HTTP 200, {ctype}, ukuran {len(body)} byte, header %PDF."
    return False, f"HTTP {response.status}, content-type {ctype or '-'}, magic PDF={magic}."


def ensure_mitigation(page) -> str:
    page.goto(f"{BASE}/mitigation/sit-nocoord", wait_until="domcontentloaded")
    link = page.locator("a[href*='/mitigation/sit-nocoord/']").first
    if link.count() == 0:
        link = page.locator("a:has-text('Print Report')").first
    if link.count() == 0:
        login(page, ADMIN)
        page.goto(f"{BASE}/mitigation/sit-nocoord", wait_until="domcontentloaded")
        page.fill("textarea[name=slope_condition]", "Cracks on the crest")
        page.fill("textarea[name=mitigation_strategy]", "Soil nailing grid")
        page.locator("button:has-text('Save Estimation Plan')").click()
        page.wait_for_load_state("domcontentloaded")
        login(page, VIEWER)
        page.goto(f"{BASE}/mitigation/sit-nocoord", wait_until="domcontentloaded")
        link = page.locator("a:has-text('Print Report')").first
    href = link.get_attribute("href") or ""
    return href.rstrip("/").split("/")[-1]


def run(page) -> dict[int, dict]:
    results: dict[int, dict] = {}
    login(page, VIEWER)

    page.goto(f"{BASE}/inventory/sit-high", wait_until="domcontentloaded")
    highlight_target(page, "a[href*='pdf-inventory']", "Print Report Inventory")
    cut = page.request.get(f"{BASE}/pdf-inventory/sit-high")
    rock = page.request.get(f"{BASE}/pdf-inventory/sit-low")
    cut_ok, cut_note = pdf_ok(cut)
    rock_ok, rock_note = pdf_ok(rock)
    ok = cut_ok and rock_ok
    results[35] = {
        "status": "Pass" if ok else "Fail",
        "actual": (
            f"Print Report Inventory terbuka. Cut sit-high: {cut_note} Rock sit-low: {rock_note}"
            if ok
            else f"Print Report Inventory gagal. Cut sit-high: {cut_note} Rock sit-low: {rock_note}"
        ),
        "shot": shot(page, 35),
    }

    page.goto(f"{BASE}/priority/sit-high", wait_until="domcontentloaded")
    highlight_target(page, "a[href*='pdf-rating']", "Print Report Rating")
    rating = page.request.get(f"{BASE}/pdf-rating/sit-high")
    rating_ok, rating_note = pdf_ok(rating)
    results[62] = {
        "status": "Pass" if rating_ok else "Fail",
        "actual": (
            f"Print Report rating terbuka. {rating_note}"
            if rating_ok
            else f"Print Report rating gagal. {rating_note}"
        ),
        "shot": shot(page, 62),
    }

    mid = ensure_mitigation(page)
    page.goto(f"{BASE}/mitigation/sit-nocoord/{mid}", wait_until="domcontentloaded")
    highlight_target(page, "a[href*='pdf-mitigation']", "Print Report Mitigation")
    mitigation = page.request.get(f"{BASE}/pdf-mitigation/{mid}")
    mit_ok, mit_note = pdf_ok(mitigation)
    missing = page.request.get(f"{BASE}/pdf-mitigation/999999")
    missing_ok = missing.status == 404
    ok = mit_ok and missing_ok and mid.isdigit()
    results[120] = {
        "status": "Pass" if ok else "Fail",
        "actual": (
            f"Print Report mitigasi id {mid} terbuka. {mit_note} ID tidak ada: HTTP {missing.status}."
            if ok
            else f"Print Report mitigasi gagal. id={mid}. {mit_note} ID tidak ada: HTTP {missing.status}."
        ),
        "shot": shot(page, 120),
    }
    return results


def write_rows(results: dict[int, dict]) -> None:
    wb = load_workbook(XLSX)
    ws = wb["Test Case (SKN)"]
    for skn, data in results.items():
        row = 3 + skn
        points = split_points(ws.cell(row, 9).value)
        note = data["actual"]
        lines = [f"{i}. [{'Sesuai' if data['status'] == 'Pass' else 'Tidak terbukti'}] {p}" for i, p in enumerate(points, 1)]
        lines += ["", f"Catatan: {note}"]
        actual = "\n".join(lines)
        ws.cell(row, 11, data["status"])
        cell = ws.cell(row, 12, actual)
        cell.alignment = WRAP
        cell.font = CELL_FONT
        ws.cell(row, 13, TESTER)
        ws.cell(row, 14, TODAY)
        raw = ensure_raw(skn, Path(data["shot"]))
        card = EVIDENCE / f"SKN-{skn:03d}.png"
        compose_card(
            card,
            skn=skn,
            status=data["status"],
            module=str(ws.cell(row, 3).value or ""),
            field=str(ws.cell(row, 4).value or ""),
            points=points,
            note=note,
            raw=raw,
        )
        evid = ws.cell(row, 15, f"SKN-{skn:03d}.png\n{len(points)} poin · {data['status']}")
        evid.hyperlink = f"evidence/SKN-{skn:03d}.png"
        evid.font = LINK_FONT
        evid.alignment = WRAP
        ws.row_dimensions[row].height = min(200, max(48, 14 * min(actual.count("\n") + 2, 12) + 10))
        print(f"SKN-{skn:03d} {data['status']} {note}")
    wb.save(XLSX)


def main() -> None:
    RAW.mkdir(parents=True, exist_ok=True)
    EVIDENCE.mkdir(parents=True, exist_ok=True)
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        context = browser.new_context(viewport={"width": 1400, "height": 900})
        page = context.new_page()
        page.set_default_timeout(20000)
        try:
            results = run(page)
        finally:
            context.close()
            browser.close()
    write_rows(results)
    failed = [k for k, v in results.items() if v["status"] != "Pass"]
    if failed:
        raise SystemExit(f"Fail SKN: {failed}")


if __name__ == "__main__":
    main()
