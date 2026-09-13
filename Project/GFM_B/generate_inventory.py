"""Workbook test case: GFM_B — modul Inventory (Dashboard Slope GFM)."""
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parent.parent
sys.path.insert(0, str(REPO))

from xlsx_kit import (  # noqa: E402
    CELL_FONT,
    GREEN,
    SKN_HEADERS,
    SIT_HEADERS,
    TITLE_FONT,
    WRAP,
    autosize,
    style_header,
    write_rows,
)
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

OUT = HERE / "GFM_B_Testcase.xlsx"
MOD = "Inventory"

ADMIN = (
    "User login GFM dengan role Admin (inventory view + create + delete). "
    "Sesi role aktif sudah dipilih."
)
VIEWER = (
    "User login GFM dengan role Viewer (inventory view saja, tanpa create/edit/delete). "
    "Sesi role aktif sudah dipilih."
)
NO_ACCESS = (
    "User login GFM tanpa permission inventory view "
    "(role tanpa module inventory, atau tanpa role)."
)
GUEST = "User belum login (tamu)."


def c(field, tipe, desc, pre, steps, exp, pri="High"):
    return (field, tipe, desc, pre, steps, exp, pri)


CASES = [
    c(
        "Guest - Redirect Login",
        "Negative",
        "Tamu tidak bisa membuka Inventory; diarahkan ke login.",
        GUEST,
        "1. Buka URL /inventory tanpa sesi login.\n"
        "2. Buka /inventory/{slug} milik slope yang sudah ada (jika diketahui).",
        "Keduanya redirect ke /login.\n"
        "Bukan 403, bukan halaman inventory.",
        "High",
    ),
    c(
        "Index - Page Load Admin",
        "Positive",
        "Halaman daftar Inventory tampil lengkap untuk Admin.",
        ADMIN + " Buka sidebar Inventory (/inventory).",
        "1. Klik menu Inventory di sidebar.\n"
        "2. Tunggu halaman selesai load.",
        "Breadcrumb: Home / Slope Inventory.\n"
        "Judul: Slope Inventory.\n"
        "Empat kartu: Total Slopes, High Risk (merah), Medium Risk (amber), Low Risk (hijau).\n"
        "Tabel kolom: Slope Name, Type, Side, STA Range, Location, RS Score, Risk, Created At, Actions.\n"
        "Tombol Add New Slope tampil.\n"
        "Setiap baris punya tombol Details.",
        "High",
    ),
    c(
        "Index - Viewer No Add",
        "Negative",
        "Viewer melihat daftar tapi tidak bisa menambah slope.",
        VIEWER + " Buka /inventory.",
        "1. Buka Inventory.\n"
        "2. Cari tombol Add New Slope dan modal Add Slope Record.",
        "List dan kartu statistik tampil.\n"
        "Tombol Add New Slope tidak dirender.\n"
        "Modal Add tidak ada di halaman.",
        "High",
    ),
    c(
        "Index - No Module Access 403",
        "Negative",
        "Tanpa inventory view, URL Inventory ditolak 403.",
        NO_ACCESS,
        "1. Buka /inventory.\n"
        "2. Buka /pdf-inventory/{slug} jika slug diketahui.",
        "HTTP 403 Forbidden (bukan redirect dashboard, bukan list kosong).",
        "High",
    ),
    c(
        "Index - Empty State",
        "Positive",
        "Tabel tanpa slope menampilkan empty state sesuai role.",
        ADMIN + " Database slopes kosong (atau filter 0 row jika filter ada).",
        "1. Buka /inventory saat belum ada slope.",
        "Satu baris: No slopes yet. Use Add New Slope to create one.\n"
        "Kartu Total = 0, High/Medium/Low = 0.\n"
        "Viewer (ulangi dengan role Viewer): teks No slopes yet. tanpa ajakan Add.",
        "Medium",
    ),
    c(
        "Index - Stats vs Stored RS",
        "Positive",
        "Kartu High/Medium/Low dihitung dari ranking RS yang tersimpan, bukan rumus JS berbeda.",
        ADMIN + " Ada slope High (RS>=1000), Medium (RS>=100), Low, dan satu tanpa ranking.",
        "1. Buka /inventory.\n"
        "2. Bandingkan hitungan kartu dengan badge Risk di tabel dan nilai RS.",
        "High: RS >= 1000, badge merah.\n"
        "Medium: RS >= 100 dan < 1000, badge amber.\n"
        "Low: selain itu atau RS kosong (tampilan RS —), badge hijau.\n"
        "Total Slopes = jumlah baris.\n"
        "Type di tabel: cut-type tampil cut (tanpa suffix -type).",
        "High",
    ),
    c(
        "Index - Column Filter",
        "Positive",
        "Filter per kolom (teks + dropdown Type/Side) sesuai BRD.",
        ADMIN + " Minimal beberapa slope beda nama, type, sisi.",
        "1. Amati header tabel: ada input filter teks dan dropdown Type/Side.\n"
        "2. Filter nama sebagian. Amati baris.\n"
        "3. Filter Type Cut. Filter Side A.\n"
        "4. Kosongkan filter.",
        "Copy UI: Use column filters to narrow results.\n"
        "Teks memfilter kolom terkait. Type/Side dropdown memfilter exact.\n"
        "Kombinasi AND. Reset mengembalikan semua baris.\n"
        "Jika input filter tidak ada sama sekali: Fail (gap vs BRD 02-inventory).",
        "High",
    ),
    c(
        "Index - Open Details",
        "Positive",
        "Tombol Details membuka halaman detail slope.",
        ADMIN + " Minimal 1 slope di list.",
        "1. Klik Details pada satu baris.",
        "Navigasi ke /inventory/{slug}.\n"
        "Header Feature No = Slope Name baris itu.",
        "High",
    ),
    c(
        "Add Modal - Open Close",
        "Positive",
        "Add New Slope membuka modal; Cancel/X menutup tanpa simpan.",
        ADMIN + " Buka /inventory.",
        "1. Klik Add New Slope.\n"
        "2. Isi Slope Name sembarang. Klik Cancel.\n"
        "3. Buka lagi, klik × di header.",
        "Modal judul: Add Slope Record.\n"
        "Subjudul: Step 1 of 3: Basic Location Details.\n"
        "Cancel dan × menutup modal. Data tidak masuk tabel.",
        "Medium",
    ),
    c(
        "Add Modal - Fields and Types",
        "Positive",
        "Form awal hanya data lokasi dasar; type dan side sesuai v1.",
        ADMIN + " Modal Add terbuka.",
        "1. Amati field dan opsi dropdown.",
        "Field: Slope Name, Location, STA 1, STA 2, Latitude, Longitude, Side of Road, Slope Type.\n"
        "Side: A, B.\n"
        "Type: Cut, Rock, Fill, Retaining, Combine "
        "(value cut-type / rock-type / fill-type / retaining-type / combine-type).\n"
        "Bukan satu form panjang geometry+rating.",
        "High",
    ),
    c(
        "Add - Slope Name Mandatory",
        "Negative",
        "Slope Name wajib.",
        ADMIN + " Modal Add terbuka.",
        "1. Kosongkan Slope Name.\n"
        "2. Isi Latitude dan Longitude, pilih Side dan Type.\n"
        "3. Klik Save & Continue.",
        "Tidak tersimpan. Browser required dan/atau validasi Laravel.\n"
        "Tetap di form/modal.",
        "High",
    ),
    c(
        "Add - Latitude Longitude Mandatory",
        "Negative",
        "Latitude dan Longitude wajib.",
        ADMIN + " Modal Add terbuka.",
        "1. Isi Slope Name unik. Kosongkan Latitude atau Longitude.\n"
        "2. Save & Continue.",
        "Tidak tersimpan. Pesan validasi / required.\n"
        "Tidak lanjut ke wizard Geometry.",
        "High",
    ),
    c(
        "Add - Duplicate Slope Name",
        "Negative",
        "Slope Name harus unik.",
        ADMIN + " Sudah ada slope bernama SP-DUP.",
        "1. Add New Slope.\n"
        "2. Slope Name = SP-DUP. Isi lat/long, Side, Type.\n"
        "3. Save & Continue.",
        "Gagal simpan. Pesan error (unique). Modal tetap terbuka (errors any).\n"
        "Tidak ada baris duplikat.",
        "High",
    ),
    c(
        "Add - Viewer Cannot POST",
        "Negative",
        "Viewer tidak bisa create lewat POST langsung.",
        VIEWER,
        "1. Kirim POST /inventory dengan payload valid (DevTools/Postman) "
        "atau pastikan tidak ada UI create.",
        "HTTP 403. Tidak ada row baru di tabel slopes.",
        "High",
    ),
    c(
        "Add - Save Continue Geometry",
        "Positive",
        "Simpan lokasi dasar lalu lanjut wizard Geometry.",
        ADMIN,
        "1. Add New Slope.\n"
        "2. Slope Name unik (contoh INV-TC-01), Location KM 12, STA 0+000 / 0+100, "
        "Latitude -7.1, Longitude 110.4, Side A, Type Cut.\n"
        "3. Save & Continue.",
        "Redirect /create/geometry/{slug}. Slug dari nama (contoh inv-tc-01).\n"
        "Flash: Slope record saved. Continue with geometry.\n"
        "Row slopes terbuat. Judul halaman Geometry — {nama}.",
        "High",
    ),
    c(
        "Wizard - Geometry Save",
        "Positive",
        "Geometry step: feature height wajib, lanjut Characteristic.",
        ADMIN + " Baru saja save slope, berada di /create/geometry/{slug}.",
        "1. Isi Feature height (m) = 10, Notes opsional.\n"
        "2. Continue to Characteristic.",
        "Redirect /create/characteristic/{slug}.\n"
        "Geometry tersimpan di session (belum wajib ke DB sampai rating disimpan).",
        "High",
    ),
    c(
        "Wizard - Characteristic Save",
        "Positive",
        "Characteristic: lima kelompok v1, lanjut Rating.",
        ADMIN + " Berada di /create/characteristic/{slug} setelah geometry.",
        "1. Isi Slope Protection, Surface Drainage, Hydrogeological Settings, "
        "Geological Features, Signs of Distress (boleh teks singkat).\n"
        "2. Continue to Rating.",
        "Redirect /create/rating/{slug}.\n"
        "Field nama kelompok sama v1 (jangan diganti).",
        "High",
    ),
    c(
        "Wizard - Rating Cut Type RS",
        "Positive",
        "Save rating Cut menghitung IS/CS/TS/RS di backend dan isi jadwal inspeksi.",
        ADMIN + " Slope type cut-type. Session geometry feature_height=10. "
        "Characteristic sudah diisi. Halaman Rating terbuka.",
        "1. Amati field A1–A5, B1, B2, C1, C2, D1, D2 dan Consequence to life.\n"
        "2. Biarkan faktor = 1. Consequence Category 2.\n"
        "3. Save parameters.",
        "Redirect detail /inventory/{slug}. Flash: Slope parameters saved successfully.\n"
        "Untuk faktor 1 dan height 10 (cut): IS=1, CS=20, TS=20, RS=1.26 (TS×0.063).\n"
        "Risk Low. Scorecard kanan menampilkan RS/IS/CS/TS tersimpan (bukan hitung ulang JS).\n"
        "Jadwal Engineer +5 tahun, Maintenance +1 tahun (category 1/2).",
        "High",
    ),
    c(
        "Wizard - Rating Category 3 Schedule",
        "Positive",
        "Consequence category-3: inspeksi Engineer 10 tahun, Maintenance 2 tahun.",
        ADMIN + " Wizard rating slope cut-type, geometry sudah di session.",
        "1. Isi faktor valid. Consequence to life = Category 3.\n"
        "2. Save parameters.\n"
        "3. Cek Inspection Schedule di detail.",
        "Engineer = hari ini + 10 tahun. Maintenance = hari ini + 2 tahun.\n"
        "Format tampilan d M Y.",
        "High",
    ),
    c(
        "Wizard - Rating Rock Scale of Failure",
        "Positive",
        "Rock-type menampilkan Scale of failure; rumus CS memakai k Small/Medium/Large.",
        ADMIN + " Slope type rock-type, sampai halaman Rating.",
        "1. Pastikan ada dropdown Scale of failure: Small, Medium, Large.\n"
        "2. Pilih Large. Faktor 1. Save.\n"
        "3. Bandingkan RS di scorecard dengan rumus rock (k=5, RS=TS×0.022).",
        "Field Scale of failure hanya untuk rock-type (tidak di cut).\n"
        "Nilai ranking tersimpan. Risk sesuai RS.\n"
        "Jangan ubah rumus; Fail jika angka menyimpang dari service backend.",
        "High",
    ),
    c(
        "Wizard - Rating Fill Extra Factors",
        "Positive",
        "Fill-type memakai faktor tambahan C21–C23 dan D21–D23.",
        ADMIN + " Slope type fill-type, halaman Rating.",
        "1. Amati field extra C21, C22, C23, D21, D22, D23.\n"
        "2. Isi 1, Save parameters.",
        "Field extra tampil hanya fill-type.\n"
        "IS/CS di scorecard boleh tampil - (fill memakai IS1–3 / CS1–3); TS dan RS terisi.\n"
        "RS fill = TS × 0.006.",
        "High",
    ),
    c(
        "Detail - Header and Tabs",
        "Positive",
        "Detail: header, aksi, tab kiri, scorecard kanan.",
        ADMIN + " Slope sudah punya ranking. Buka Details.",
        "1. Amati header dan tombol aksi.\n"
        "2. Klik tab Basic Information, Detail Information, Photos, Documents, Records.",
        "Header: Feature No, lokasi, badge {risk} risk, Side, tipe (CUT SLOPE dll), angka RS.\n"
        "Aksi Admin: Back, Print Report, Edit Configuration, Delete.\n"
        "Tab kiri berganti konten tanpa pindah URL.\n"
        "Basic: Location, STA, Latitude, Longitude.\n"
        "Detail: Geometry JSON + Characteristic JSON.\n"
        "Kanan: Hazard Scorecard (RS, IS, CS, TS) + Inspection Schedule.",
        "High",
    ),
    c(
        "Detail - Viewer Read Only",
        "Negative",
        "Viewer detail tanpa kontrol tulis.",
        VIEWER + " Slope ada. Buka /inventory/{slug}.",
        "1. Amati tombol header dan tiap tab.",
        "Back dan Print Report ada.\n"
        "Edit Configuration, Delete, Upload (foto/dokumen), Add Record tidak tampil.\n"
        "Teks Delete pada foto tidak ada.",
        "High",
    ),
    c(
        "Detail - Map Coordinates",
        "Positive",
        "Kolom kanan menampilkan peta koordinat; tanpa koordinat ada placeholder.",
        ADMIN + " Satu slope dengan lat/long terisi, satu tanpa (jika bisa).",
        "1. Buka detail yang punya Latitude/Longitude.\n"
        "2. Cari peta di kolom kanan.\n"
        "3. Jika ada slope tanpa koordinat: buka dan amati placeholder.",
        "BRD: peta koordinat di kanan. Ada koordinat: peta/marker.\n"
        "Tanpa koordinat: No coordinates available.\n"
        "Jika peta tidak ada sama sekali: Fail (gap vs BRD 02-inventory).",
        "Medium",
    ),
    c(
        "Detail - Back",
        "Positive",
        "Back kembali ke daftar Inventory.",
        ADMIN + " Halaman detail terbuka.",
        "1. Klik Back.",
        "Kembali /inventory. List tampil. Tidak ada dialog wajib (implementasi saat ini langsung Back).",
        "Medium",
    ),
    c(
        "Edit Configuration",
        "Positive",
        "Admin mengedit konfigurasi lewat wizard geometry lagi; nilai lama terisi.",
        ADMIN + " Slope sudah punya geometry/characteristic/rating.",
        "1. Di detail, klik Edit Configuration.\n"
        "2. Pastikan Feature height terisi nilai lama.\n"
        "3. Ubah height, lanjut Characteristic (nilai lama), Rating, Save.",
        "URL /create/geometry/{slug}. Field terisi dari JSON slope.\n"
        "Setelah save, detail menampilkan parameter baru dan RS dihitung ulang.",
        "High",
    ),
    c(
        "Photos - Groups and Upload",
        "Positive",
        "Upload foto per kelompok v1, lalu tampil di tab Photos.",
        ADMIN + " Detail slope, tab Photos.",
        "1. Pastikan Category: Plan Section, Cross Section, Slope Protection, "
        "Surface Drainage Provision, Hydrogeological Settings, Geological Features, "
        "Signs of Distress, Instability.\n"
        "2. Pilih Plan Section, file .jpg valid, Upload.",
        "Flash: Photo uploaded.\n"
        "Thumbnail tampil, caption Plan Section.\n"
        "Nama kelompok tidak diubah dari v1.",
        "High",
    ),
    c(
        "Photos - Delete",
        "Positive",
        "Admin menghapus foto; file terkait ikut hilang dari tampilan.",
        ADMIN + " Slope punya minimal 1 foto.",
        "1. Tab Photos. Klik Delete pada satu foto.",
        "Flash: Image deleted successfully. Foto hilang dari grid.\n"
        "Jika semua terhapus: No photos yet.",
        "High",
    ),
    c(
        "Photos - Invalid Type",
        "Negative",
        "Upload non-image ditolak.",
        ADMIN + " Tab Photos.",
        "1. Pilih Category Plan Section.\n"
        "2. Pilih file .pdf atau .txt. Upload.",
        "Validasi gagal (image). Foto tidak masuk daftar.\n"
        "accept=image/* di input file.",
        "High",
    ),
    c(
        "Photos - Max Size",
        "Negative",
        "Foto lebih dari 20480 KB (20 MB) ditolak.",
        ADMIN + " Siapkan file gambar > 20 MB.",
        "1. Tab Photos. Upload file gambar > 20 MB.",
        "Validasi max 20480 gagal. Tidak ada Photo uploaded.\n"
        "Daftar foto tidak bertambah.",
        "Medium",
    ),
    c(
        "Photos - Viewer Cannot Upload",
        "Negative",
        "Viewer tidak bisa POST foto.",
        VIEWER + " Slope ada.",
        "1. Buka detail, tab Photos: tidak ada form Upload.\n"
        "2. (Opsional) POST /inventory/{slug}/photos sebagai viewer.",
        "Tidak ada kontrol Upload. POST langsung → 403.",
        "High",
    ),
    c(
        "Documents - Upload and List",
        "Positive",
        "Admin upload dokumen; nama asli tampil di tab Documents.",
        ADMIN + " Tab Documents.",
        "1. Pilih file PDF/Word/Excel valid, Upload.\n"
        "2. Amati daftar.",
        "Flash: Document uploaded.\n"
        "List menampilkan original_name.\n"
        "Tab label Documents (n) bertambah.\n"
        "File fisik di storage, bukan BLOB DB.",
        "High",
    ),
    c(
        "Documents - Empty",
        "Positive",
        "Tanpa dokumen: empty copy.",
        ADMIN + " Slope tanpa dokumen.",
        "1. Buka tab Documents.",
        "Teks: No documents. Form Upload tetap ada untuk Admin.",
        "Low",
    ),
    c(
        "Records - Add",
        "Positive",
        "Admin menambah catatan kejadian di tab Records.",
        ADMIN + " Tab Records.",
        "1. Isi Date of occurrence dan Type of failure.\n"
        "2. Add Record.",
        "Flash: Record saved.\n"
        "Baris baru tampil. Tab Records (n) bertambah.\n"
        "Kosong: No records.",
        "High",
    ),
    c(
        "PDF - Print Report Viewer",
        "Positive",
        "Viewer dan Admin bisa unduh PDF inventory.",
        VIEWER + " Slope ada (cut atau rock).",
        "1. Di detail, klik Print Report.\n"
        "2. Ulangi sebagai Admin.",
        "Response PDF (Content-Type application/pdf).\n"
        "Nama file inventory-{slug}.pdf (inline).\n"
        "Rock-type memakai template rock; selain itu cut_inventory.\n"
        "Bukan 403 untuk user ber-inventory view.",
        "High",
    ),
    c(
        "PDF - Unknown Slug",
        "Negative",
        "Slug tidak ada → 404.",
        ADMIN,
        "1. Buka /pdf-inventory/slug-tidak-ada.\n"
        "2. Buka /inventory/slug-tidak-ada.",
        "HTTP 404. Bukan 500.",
        "Medium",
    ),
    c(
        "Delete - Confirm and Cascade",
        "Positive",
        "Hapus slope minta konfirmasi lalu hapus data terkait.",
        ADMIN + " Slope tes dengan foto/dokumen/record (boleh minimal).",
        "1. Di detail, klik Delete. Batalkan confirm.\n"
        "2. Klik Delete lagi, OK pada Delete this slope?",
        "Batal: tetap di detail.\n"
        "Ya: redirect /inventory, flash Item deleted successfully.\n"
        "Slope hilang dari list. Record, dokumentasi, inspeksi/mitigasi/preservasi terkait (jika ada) ikut terhapus. Folder storage slope ikut dibersihkan.",
        "High",
    ),
    c(
        "Delete - Viewer Forbidden",
        "Negative",
        "Viewer tidak bisa hapus slope.",
        VIEWER + " Slope ada.",
        "1. Detail: pastikan tombol Delete tidak ada.\n"
        "2. DELETE /slope/{slug} sebagai viewer.",
        "Tidak ada tombol Delete. Request langsung 403. Slope tetap ada.",
        "High",
    ),
]


def main():
    from generate_gfm_b import main as generate_all

    generate_all()
    return

    wb = Workbook()

    ws = wb.active
    ws.title = "Test Case (SKN)"
    ws["A1"] = "TEST CASE — DASHBOARD SLOPE GFM — INVENTORY"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(SKN_HEADERS))
    ws["A2"] = (
        "Sumber: docs/brd/02-inventory.md, InventoryController, views inventory/*, "
        "CalculateRankingService. Folder Project/GFM_B. "
        "Kolom Status Uji / Hasil Aktual / QA / Tanggal Uji diisi saat eksekusi. "
        "Sheet Issue SIT: isi baris baru bila SKN Fail; Ref ID = Scenario ID di sheet ini."
    )
    ws["A2"].alignment = WRAP
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(SKN_HEADERS))
    ws.row_dimensions[2].height = 48

    for col, header in enumerate(SKN_HEADERS, 1):
        ws.cell(3, col, header)
    style_header(ws, 3, len(SKN_HEADERS))

    skn_rows = []
    for i, (field, tipe, desc, pre, steps, exp, pri) in enumerate(CASES, 1):
        skn_rows.append(
            [
                i,
                f"SKN-{i:03d}",
                MOD,
                field,
                tipe,
                desc,
                pre,
                steps,
                exp,
                pri,
                "Not Run",
                "",
                "",
                "",
            ]
        )
    write_rows(ws, 4, skn_rows)
    for row in range(4, 4 + len(skn_rows)):
        ws.row_dimensions[row].height = 78
    ws.auto_filter.ref = f"A3:N{3 + len(skn_rows)}"
    ws.freeze_panes = "A4"
    ws.row_dimensions[3].height = 22
    autosize(ws, [6, 14, 14, 36, 12, 44, 40, 52, 52, 12, 12, 22, 14, 16])

    dv_status = DataValidation(type="list", formula1='"Not Run,Pass,Fail"', allow_blank=True)
    dv_status.add(f"K4:K{3 + len(skn_rows)}")
    ws.add_data_validation(dv_status)
    dv_tipe = DataValidation(type="list", formula1='"Positive,Negative"', allow_blank=False)
    dv_tipe.add(f"E4:E{3 + len(skn_rows)}")
    ws.add_data_validation(dv_tipe)

    sit = wb.create_sheet("Issue SIT")
    sit["A1"] = "LOG ISSUE / DEFECT SYSTEM INTEGRATION TESTING (SIT)"
    sit["A1"].font = TITLE_FONT
    sit.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(SIT_HEADERS))
    sit["A2"] = (
        "Belum ada log SIT Inventory. Tambah baris saat SKN berstatus Fail. "
        "Ref ID mengarah ke Scenario ID (SKN-xxx) di sheet Test Case."
    )
    sit["A2"].alignment = WRAP
    sit.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(SIT_HEADERS))
    sit.row_dimensions[2].height = 32
    for col, header in enumerate(SIT_HEADERS, 1):
        sit.cell(3, col, header)
    style_header(sit, 3, len(SIT_HEADERS))
    sit.freeze_panes = "A4"
    sit.auto_filter.ref = "A3:N3"
    autosize(sit, [6, 12, 12, 16, 32, 40, 42, 48, 12, 16, 14, 16, 20, 16])

    pos = sum(1 for row in skn_rows if row[4] == "Positive")
    neg = sum(1 for row in skn_rows if row[4] == "Negative")
    high = sum(1 for row in skn_rows if row[9] == "High")
    med = sum(1 for row in skn_rows if row[9] == "Medium")
    low = sum(1 for row in skn_rows if row[9] == "Low")

    summary = wb.create_sheet("Summary", 0)
    summary["A1"] = "GFM_B — INVENTORY TEST CASE SUMMARY"
    summary["A1"].font = TITLE_FONT
    summary["A3"] = "Item"
    summary["B3"] = "Jumlah"
    style_header(summary, 3, 2)
    lines = [
        ("Modul Inventory (SKN)", len(skn_rows)),
        ("Positive", pos),
        ("Negative", neg),
        ("Priority High", high),
        ("Priority Medium", med),
        ("Priority Low", low),
        ("Issue SIT (awal)", 0),
    ]
    for i, (label, n) in enumerate(lines, 4):
        summary.cell(i, 1, label).font = CELL_FONT
        summary.cell(i, 2, n).font = CELL_FONT
    total_row = 4 + len(lines)
    summary.cell(total_row, 1, "Total SKN").font = Font(bold=True, name="Calibri")
    summary.cell(total_row, 2, len(skn_rows)).font = Font(bold=True, name="Calibri")
    summary["A3"].fill = GREEN
    summary["B3"].fill = GREEN
    autosize(summary, [32, 16])

    wb.save(OUT)
    print(f"Wrote {OUT.name}: {len(skn_rows)} SKN")
    print(f"  Positive={pos} Negative={neg} High={high} Medium={med} Low={low}")


if __name__ == "__main__":
    from generate_gfm_b import main as generate_all

    generate_all()
