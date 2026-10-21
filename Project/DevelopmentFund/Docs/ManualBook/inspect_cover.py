import docx

doc = docx.Document(r"d:\Work\Source\TestCase_Generator\Project\DevelopmentFund\Manual Book\output\User Manual - DF Overview v1.2.docx")
print("Total Sections:", len(doc.sections))
for s_idx, s in enumerate(doc.sections):
    print(f"Section {s_idx}: w={s.page_width.pt}pt, h={s.page_height.pt}pt, top={s.top_margin.pt}pt, bottom={s.bottom_margin.pt}pt, left={s.left_margin.pt}pt, right={s.right_margin.pt}pt")

for i, p in enumerate(doc.paragraphs[:35]):
    text = p.text.strip()
    xml = p._p.xml
    is_break = "w:type=\"page\"" in xml or "w:pageBreakBefore" in xml
    print(f"P{i:2d}: style={p.style.name}, sb={p.paragraph_format.space_before}, sa={p.paragraph_format.space_after}, br={is_break}, text={text[:60]!r}")
    if is_break:
        print(f"--- PAGE BREAK FOUND AT P{i} ---")
        break

print("\n--- FIRST 3 TABLES ---")
for t_idx, t in enumerate(doc.tables[:3]):
    print(f"\nTable {t_idx} (rows={len(t.rows)}, cols={len(t.columns)}):")
    for r_idx, r in enumerate(t.rows):
        cells = [c.text.strip().replace("\n", " | ") for c in r.cells]
        print(f"  R{r_idx}: {cells}")
