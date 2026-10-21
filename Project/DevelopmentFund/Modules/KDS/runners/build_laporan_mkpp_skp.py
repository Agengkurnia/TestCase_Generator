"""User-testing workbook for the 12/10/2026 MKPP DF batch: 10 docs, 2 or 3 SKP."""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent.parent  # module root
STATE = HERE / "batch_10_mkpp_skp.json"
OUT = HERE.parent / "Laporan_Uji_MKPP_DF_10_12-10-2026.xlsx"

NAVY = "1F4E79"
TEAL = "0F6E6E"
AMBER = "8A5A00"
RED = "8C2F2F"
GREEN = "1E7A46"
PALE = "F4F7FB"
PALE_GREEN = "E8F6EE"
PALE_RED = "FDECEC"
WHITE = "FFFFFF"

thin = Border(
    left=Side(style="thin", color="D0D7E2"),
    right=Side(style="thin", color="D0D7E2"),
    top=Side(style="thin", color="D0D7E2"),
    bottom=Side(style="thin", color="D0D7E2"),
)

NAMES = {
    "050311": "CATUR SENTOSA ANUGERAH, PT (BANDUNG SELATAN)",
    "040310": "CATUR SENTOSA ANUGERAH, PT (JAKARTA SELATAN)",
    "400310": "CATUR SENTOSA ANUGERAH, PT (JAKARTA BARAT)",
    "400305": "CATUR SENTOSA ANUGERAH, PT (SERANG)",
    "400306": "CATUR SENTOSA ANUGERAH, PT (TANGERANG KABUPATEN)",
    "400309": "CATUR SENTOSA ANUGERAH, PT (TANGERANG KOTA)",
}


def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)


def rupiah(value):
    return f"Rp {int(value):,}".replace(",", ".")


def period_id(us_date):
    month, day, year = us_date.split("/")
    return f"{day}/{month}/{year}"


def style_header(ws, row, start, end, color=NAVY):
    for col in range(start, end + 1):
        cell = ws.cell(row, col)
        cell.font = Font(bold=True, color=WHITE, name="Calibri", size=11)
        cell.fill = fill(color)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin


def style_body(cell, wrap=True):
    cell.font = Font(name="Calibri", size=11)
    cell.alignment = Alignment(vertical="center", wrap_text=wrap)
    cell.border = thin


def widths(ws, sizes):
    for idx, size in enumerate(sizes, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = size


def landscape(ws, title_rows):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.print_title_rows = title_rows
    ws.sheet_view.showGridLines = False


def main():
    state = json.loads(STATE.read_text(encoding="utf-8"))
    docs = sorted(state["documents"], key=lambda d: d["n"])
    period_from = period_id(state["periodFrom"])
    period_to = period_id(state["periodTo"])
    checks = {(r["docNo"], r["siteCode"]): r for r in state.get("prmChecks") or []}

    wb = Workbook()
    ws = wb.active
    ws.title = "Ringkasan MKPP"
    ws.sheet_properties.tabColor = NAVY
    ws.merge_cells("A1:L1")
    ws["A1"] = "Laporan Data Uji — 10 MKPP Development Fund (2–3 SKP)"
    ws["A1"].font = Font(bold=True, size=16, color=NAVY, name="Calibri")
    ws.merge_cells("A2:L2")
    ws["A2"] = (
        "Lingkungan DEV. Dibuat 12/10/2026 oleh Muslih, disetujui PATRICK.YOHANNES. "
        "Group CSA. Satu activity CONSUMER OFFER FEE (380202). Budget masing-masing Rp 1.000.000. "
        f"Periode {period_from} s/d {period_to}. Klaim belum dibuat."
    )
    ws["A2"].font = Font(size=11, color="333333", name="Calibri")
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 36

    headers = [
        "No",
        "Nomor MKPP",
        "Status KDS",
        "Group",
        "Tipe Budget",
        "Periode Dari",
        "Periode Sampai",
        "Total Budget",
        "Jumlah SKP",
        "Subdist",
        "Pembuat",
        "Approver",
    ]
    for col, header in enumerate(headers, start=1):
        ws.cell(4, col, header)
    style_header(ws, 4, 1, len(headers))
    ws.row_dimensions[4].height = 30
    ws.freeze_panes = "A5"

    for i, doc in enumerate(docs):
        row = 5 + i
        values = [
            doc["n"],
            doc["docNo"],
            "APPROVED",
            "CSA",
            "DEVFUND",
            period_from,
            period_to,
            rupiah(doc["amount"]),
            doc["skpCount"],
            f"{doc['subdist']} ({doc['site']})",
            "Muslih",
            "PATRICK.YOHANNES",
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row, col, value)
            style_body(cell)
            if i % 2 == 1:
                cell.fill = fill(PALE)
        ws.cell(row, 3).font = Font(bold=True, color=GREEN, name="Calibri", size=11)
        ws.cell(row, 3).alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row].height = 22
    last = 4 + len(docs)
    ws.auto_filter.ref = f"A4:L{last}"
    widths(ws, [6, 28, 16, 12, 14, 16, 18, 18, 14, 36, 16, 24])
    landscape(ws, "1:4")
    ws.oddHeader.left.text = "DEV — MKPP DF untuk uji user"

    wd = wb.create_sheet("Detail SKP")
    wd.sheet_properties.tabColor = TEAL
    detail_headers = [
        "No",
        "Nomor MKPP",
        "Site / Username PRM",
        "Subdist",
        "Activity",
        "Budget Activity",
        "No SKP 1",
        "Nominal SKP 1",
        "No SKP 2",
        "Nominal SKP 2",
        "No SKP 3",
        "Nominal SKP 3",
    ]
    for col, header in enumerate(detail_headers, start=1):
        wd.cell(1, col, header)
    style_header(wd, 1, 1, len(detail_headers), TEAL)
    wd.row_dimensions[1].height = 30
    wd.freeze_panes = "A2"
    for i, doc in enumerate(docs):
        row = 2 + i
        skp = doc["skp"]
        values = [
            doc["n"],
            doc["docNo"],
            doc["site"],
            NAMES.get(doc["site"], doc["subdist"]),
            "CONSUMER OFFER FEE (380202)",
            rupiah(doc["amount"]),
            skp[0]["no"],
            rupiah(skp[0]["amount"]),
            skp[1]["no"],
            rupiah(skp[1]["amount"]),
            skp[2]["no"] if len(skp) > 2 else "",
            rupiah(skp[2]["amount"]) if len(skp) > 2 else "",
        ]
        for col, value in enumerate(values, start=1):
            cell = wd.cell(row, col, value)
            style_body(cell, wrap=False)
            if i % 2 == 1:
                cell.fill = fill(PALE)
        wd.row_dimensions[row].height = 20
    wd.auto_filter.ref = f"A1:L{1 + len(docs)}"
    widths(wd, [6, 28, 22, 48, 32, 18, 28, 16, 28, 16, 28, 16])
    landscape(wd, "1:1")

    wp = wb.create_sheet("Cek PRM")
    wp.sheet_properties.tabColor = AMBER
    prm_headers = [
        "No",
        "Nomor MKPP",
        "Subdist",
        "Username PRM",
        "Jumlah SKP",
        "Budget",
        "Muncul di LOV Memo",
        "Status pada LOV",
        "Siap diklaim",
        "Teks LOV",
        "Catatan",
    ]
    for col, header in enumerate(prm_headers, start=1):
        wp.cell(1, col, header)
    style_header(wp, 1, 1, len(prm_headers), AMBER)
    wp.row_dimensions[1].height = 30
    wp.freeze_panes = "A2"
    ready_note = "Muncul di LOV Memo MKPP. Belum dipilih dan belum disimpan sebagai klaim."
    missing_note = "Tidak ketemu di LOV memo subdist ini. Klaim tidak dibuat."
    for i, doc in enumerate(docs):
        row = 2 + i
        hit = checks.get((doc["docNo"], doc["site"]), {})
        found = bool(hit.get("found"))
        text = hit.get("text") or ""
        status = "APPROVED" if found and "APPROVED" in text.upper() else ("-" if not found else "lihat teks LOV")
        values = [
            doc["n"],
            doc["docNo"],
            doc["subdist"],
            doc["site"],
            doc["skpCount"],
            rupiah(doc["amount"]),
            "Ya" if found else "Tidak",
            status,
            "Ya, belum diklaim" if found else "Belum",
            text,
            ready_note if found else missing_note,
        ]
        for col, value in enumerate(values, start=1):
            cell = wp.cell(row, col, value)
            style_body(cell)
            cell.fill = fill(PALE_GREEN if found else PALE_RED)
        flag = wp.cell(row, 7)
        flag.font = Font(bold=True, color=GREEN if found else RED, name="Calibri", size=11)
        flag.alignment = Alignment(horizontal="center", vertical="center")
        wp.row_dimensions[row].height = 36
    wp.auto_filter.ref = f"A1:K{1 + len(docs)}"
    widths(wp, [6, 28, 26, 18, 14, 16, 20, 18, 20, 55, 48])
    landscape(wp, "1:1")

    wu = wb.create_sheet("Cara Uji")
    wu.sheet_properties.tabColor = "5B4B8A"
    wu["A1"] = "Cara pakai data ini"
    wu["A1"].font = Font(bold=True, size=16, color=NAVY, name="Calibri")
    wu.merge_cells("A1:B1")
    steps = [
        ("Portal KDS", "https://kicaokdsdev.kalbenutritionals.web.id/Transaction/MKPPDF"),
        ("Portal PRM", "https://newprmdev.kalbenutritionals.web.id/Klaim/KlaimKDSDF"),
        ("Password DEV", "bypass"),
        ("Pembuat dokumen", "Muslih"),
        ("Approver", "PATRICK.YOHANNES"),
        ("Periode uji", f"{period_from} s/d {period_to}"),
        ("Budget per MKPP", "Rp 1.000.000, satu activity CONSUMER OFFER FEE (380202), dibagi ke 2 atau 3 SKP"),
        (
            "Pembagian SKP",
            "2 SKP: Rp 500.000 + Rp 500.000. 3 SKP: Rp 400.000 + Rp 300.000 + Rp 300.000.",
        ),
        (
            "Subdist",
            "Bandung Selatan (050311), Jakarta Selatan (040310), Jakarta Barat (400310), "
            "Serang (400305), Tangerang Kabupaten (400306), Tangerang Kota (400309). "
            "Username PRM = kode site.",
        ),
        (
            "Langkah cek tanpa klaim",
            "1. Login PRM dengan username site subdist.\n"
            "2. Buka Klaim KDS DF.\n"
            "3. Klik New.\n"
            "4. Buka LOV Memo MKPP dan cari nomor dokumen.\n"
            "5. Pastikan status APPROVED dan outstanding masih Rp 1.000.000.\n"
            "6. Tutup LOV tanpa klik Pilih dan tanpa Save.",
        ),
        (
            "Yang diuji user",
            "Pilih memo, isi klaim sesuai budget activity Rp 1.000.000 dan jumlah SKP pada dokumen, lalu simpan klaim sesuai skenario uji.",
        ),
    ]
    wu["A3"] = "Item"
    wu["B3"] = "Keterangan"
    style_header(wu, 3, 1, 2, "5B4B8A")
    for i, (label, text) in enumerate(steps):
        row = 4 + i
        a = wu.cell(row, 1, label)
        b = wu.cell(row, 2, text)
        style_body(a)
        style_body(b)
        a.font = Font(bold=True, name="Calibri", size=11)
        a.fill = fill(PALE)
        wu.row_dimensions[row].height = 32 if "\n" not in text else 20 + 16 * text.count("\n")
    widths(wu, [28, 110])
    landscape(wu, "1:3")
    wu.oddFooter.center.text = "Data uji DEV 12/10/2026 — jangan pakai di production"

    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
