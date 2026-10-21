"""Generator Excel Workbook Standar (Summary, Test Case SKN, Issue SIT, Performance Metrics)."""
from __future__ import annotations

import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parent
sys.path.insert(0, str(REPO))
sys.path.insert(0, str(HERE))

from cases_template import CASES, MOD  # noqa: E402
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation
from xlsx_kit import (  # noqa: E402
    CELL_FONT,
    GREEN,
    SKN_HEADERS,
    SIT_HEADERS,
    TITLE_FONT,
    WRAP,
    autosize,
    style_header,
    write_rows,
)

OUT = HERE / "KDS_DF_Testcase.xlsx"

MODULES = [
    (MOD, CASES),
]


def packed():
    rows = []
    for mod, cases in MODULES:
        for field, tipe, desc, pre, steps, exp, pri in cases:
            rows.append((mod, field, tipe, desc, pre, steps, exp, pri))
    return rows


def main():
    skn = packed()
    wb = Workbook()

    # Sheet 1: Test Case (SKN)
    ws = wb.active
    ws.title = "Test Case (SKN)"
    ws["A1"] = f"TEST CASE — {MOD.upper()} (FUNCTIONAL & PERFORMANCE SLA)"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(SKN_HEADERS))
    ws["A2"] = (
        "Mencakup pengujian Fungsional, Pop-up Modal, dan SLA Kinerja. "
        "Kolom Status Uji / Hasil Aktual / Bukti Pengujian / QA / Tanggal Uji diisi saat eksekusi Playwright."
    )
    ws["A2"].alignment = WRAP
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(SKN_HEADERS))
    ws.row_dimensions[2].height = 48

    for col, header in enumerate(SKN_HEADERS, 1):
        ws.cell(3, col, header)
    style_header(ws, 3, len(SKN_HEADERS))

    skn_rows = []
    for i, (mod, field, tipe, desc, pre, steps, exp, pri) in enumerate(skn, 1):
        skn_rows.append([
            i, f"SKN-{i:03d}", mod, field, tipe, desc, pre, steps, exp, pri,
            "Not Run", "", "Lihat Evidence", "", ""
        ])
    write_rows(ws, 4, skn_rows)
    for row in range(4, 4 + len(skn_rows)):
        ws.row_dimensions[row].height = 78
    ws.auto_filter.ref = f"A3:O{3 + len(skn_rows)}"
    ws.freeze_panes = "A4"
    ws.row_dimensions[3].height = 22
    autosize(ws, [6, 14, 22, 38, 12, 44, 40, 52, 52, 12, 12, 24, 18, 14, 16])

    dv_status = DataValidation(type="list", formula1='"Not Run,Pass,Fail"', allow_blank=True)
    dv_status.add(f"K4:K{3 + len(skn_rows)}")
    ws.add_data_validation(dv_status)
    dv_tipe = DataValidation(type="list", formula1='"Positive,Negative"', allow_blank=False)
    dv_tipe.add(f"E4:E{3 + len(skn_rows)}")
    ws.add_data_validation(dv_tipe)

    # Sheet 2: Issue SIT
    sit = wb.create_sheet("Issue SIT")
    sit["A1"] = "LOG ISSUE / DEFECT SYSTEM INTEGRATION TESTING (SIT)"
    sit["A1"].font = TITLE_FONT
    sit.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(SIT_HEADERS))
    sit["A2"] = "Kosong di awal generate. Runner mengisi baris saat SKN Fail. Ref ID = Scenario ID (SKN-xxx)."
    sit["A2"].alignment = WRAP
    sit.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(SIT_HEADERS))
    sit.row_dimensions[2].height = 32
    for col, header in enumerate(SIT_HEADERS, 1):
        sit.cell(3, col, header)
    style_header(sit, 3, len(SIT_HEADERS))
    sit.freeze_panes = "A4"
    sit.auto_filter.ref = "A3:N3"
    autosize(sit, [6, 12, 12, 16, 32, 40, 42, 48, 12, 16, 14, 16, 20, 16])

    # Sheet 3: Performance Metrics
    perf_headers = [
        "No.", "Metric / Test Area", "SLA Target (ms)", "Measured Time (ms)", "Status", "Performance Evaluation & Observations"
    ]
    perf = wb.create_sheet("Performance Metrics")
    perf["A1"] = f"{MOD.upper()} — PERFORMANCE BENCHMARK & SLA REPORT"
    perf["A1"].font = TITLE_FONT
    perf.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(perf_headers))
    perf["A2"] = "Metrik kinerja diukur langsung menggunakan browser Performance API & Resource Timing."
    perf["A2"].alignment = WRAP
    perf.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(perf_headers))
    for col, header in enumerate(perf_headers, 1):
        perf.cell(3, col, header)
    style_header(perf, 3, len(perf_headers))
    perf.freeze_panes = "A4"
    autosize(perf, [6, 38, 18, 20, 14, 55])

    # Sheet 4: Summary (Front Sheet)
    counts = {}
    for row in skn_rows:
        counts[row[2]] = counts.get(row[2], 0) + 1

    summary = wb.create_sheet("Summary", 0)
    summary["A1"] = f"{MOD.upper()} — DASHBOARD TEST SUMMARY"
    summary["A1"].font = TITLE_FONT
    summary["A3"] = "Modul / Area"
    summary["B3"] = "Jumlah SKN"
    style_header(summary, 3, 2)
    r = 4
    for m_name, n in counts.items():
        summary.cell(r, 1, m_name).font = CELL_FONT
        summary.cell(r, 2, n).font = CELL_FONT
        r += 1
    summary.cell(r, 1, "Total SKN").font = Font(bold=True, name="Calibri")
    summary.cell(r, 2, len(skn_rows)).font = Font(bold=True, name="Calibri")
    summary["A3"].fill = GREEN
    summary["B3"].fill = GREEN
    autosize(summary, [32, 16])

    wb.save(OUT)
    print(f"Wrote {OUT.name}: {len(skn_rows)} SKN")
    return skn_rows


if __name__ == "__main__":
    main()
