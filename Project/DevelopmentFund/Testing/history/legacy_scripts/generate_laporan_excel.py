"""Generate human-friendly Excel laporan from FINAL/RINGKAS content."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from _pack import REPORTS, EVIDENCE_ROOT

OUT = REPORTS / "Laporan_Uji_DF_DEV_2026-10-18.xlsx"
EVIDENCE = EVIDENCE_ROOT / "evidence_dual_2026_10_18_10-02-29"

GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
ORANGE = PatternFill("solid", fgColor="FCE4D6")
BLUE = PatternFill("solid", fgColor="DDEBF7")
HEAD = PatternFill("solid", fgColor="1E7D37")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE = Font(bold=True, size=14, color="1E7D37")
BOLD = Font(bold=True, size=11)
THIN = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)
WRAP = Alignment(wrap_text=True, vertical="top")


def style_header(ws, cols: int, row: int = 1):
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.fill = HEAD
        cell.font = HEAD_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def autosize(ws, widths: list[int]):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def border_range(ws, r1, r2, c1, c2):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(r, c).border = THIN
            ws.cell(r, c).alignment = WRAP


def main():
    wb = Workbook()

    # ─── 1. Ringkas (Awam) ───
    ws = wb.active
    ws.title = "1. Ringkas (Awam)"
    ws["A1"] = "Ringkas Hasil Uji — Development Fund (DEV)"
    ws["A1"].font = TITLE
    ws.merge_cells("A1:D1")

    meta = [
        ("Tanggal", "18/10/2026"),
        ("QA", "Ageng"),
        ("Lingkungan", "DEV (mavendev / kicaokdsdev / newprmdev)"),
        ("Batas nominal uji", "Tiap kontrak MKPP maksimal Rp 100.000"),
        ("Status akhir", "Berhasil sebagian — belum selesai end-to-end"),
    ]
    ws["A3"] = "Informasi"
    ws["A3"].font = BOLD
    ws.append([])  # row 4 unused after write below
    r = 4
    for k, v in meta:
        ws.cell(r, 1, k).font = BOLD
        ws.cell(r, 2, v)
        if k == "Status akhir":
            ws.cell(r, 2).fill = ORANGE
            ws.cell(r, 2).font = BOLD
        r += 1

    ws.cell(r + 1, 1, "Dalam satu kalimat").font = BOLD
    ws.merge_cells(start_row=r + 2, start_column=1, end_row=r + 2, end_column=4)
    ws.cell(
        r + 2,
        1,
        "Sistem sudah bisa buat kontrak → disetujui → buat klaim → ditandai siap kirim, "
        "tetapi pengiriman klaim terakhir masih gagal, jadi proses klaim belum sampai status final.",
    )
    ws.cell(r + 2, 1).alignment = WRAP
    ws.row_dimensions[r + 2].height = 45

    start = r + 4
    ws.cell(start, 1, "Alur 5 langkah").font = BOLD
    headers = ["No", "Langkah", "Hasil", "Keterangan"]
    for i, h in enumerate(headers, 1):
        ws.cell(start + 1, i, h)
    style_header(ws, 4, start + 1)
    steps = [
        (1, "Buat kontrak dana (MKPP)", "Berhasil", ""),
        (2, "Atasan setujui kontrak", "Berhasil", "Patrick Yohannes (satu tingkat)"),
        (3, "Subdist buat & simpan klaim", "Berhasil", "Portal PRM"),
        (4, 'ABM tandai "siap kirim"', "Berhasil", "Fahmi Permadi"),
        (5, "Kirim klaim ke tahap akhir", "Gagal", "Putus di sini — matching"),
    ]
    for i, row in enumerate(steps):
        rr = start + 2 + i
        for c, val in enumerate(row, 1):
            ws.cell(rr, c, val)
        fill = GREEN if row[2] == "Berhasil" else RED
        ws.cell(rr, 3).fill = fill
        ws.cell(rr, 3).font = BOLD
    border_range(ws, start + 1, start + 6, 1, 4)

    fix_start = start + 8
    ws.cell(fix_start, 1, "Yang perlu diperbaiki dulu").font = BOLD
    for i, h in enumerate(["Prioritas", "Masalah (bahasa biasa)", "Dampak ke user"], 1):
        ws.cell(fix_start + 1, i, h)
    style_header(ws, 3, fix_start + 1)
    fixes = [
        ("Paling penting", 'Saat kirim klaim, sistem menolak karena "total matching" tidak sama dengan total klaim', "Klaim macet di draft; proses tidak selesai"),
        ("Penting", "User masih bisa lanjut simpan/submit meski data belum lengkap", "Risiko dokumen kosong / salah tersimpan"),
        ("Penting", "Petugas ABM sulit mencari klaim lewat kotak pencarian", "Ribet untuk verifikasi harian"),
    ]
    for i, row in enumerate(fixes):
        rr = fix_start + 2 + i
        for c, val in enumerate(row, 1):
            ws.cell(rr, c, val)
        ws.cell(rr, 1).fill = RED if i == 0 else YELLOW
    border_range(ws, fix_start + 1, fix_start + 4, 1, 3)

    dec = fix_start + 6
    ws.cell(dec, 1, "Keputusan").font = BOLD
    ws.merge_cells(start_row=dec + 1, start_column=1, end_row=dec + 1, end_column=4)
    ws.cell(
        dec + 1,
        1,
        "Sampai tiga poin perbaikan di atas beres, uji end-to-end belum bisa dinyatakan lulus penuh.",
    )
    ws.cell(dec + 1, 1).fill = ORANGE
    ws.cell(dec + 1, 1).alignment = WRAP
    ws.row_dimensions[dec + 1].height = 35

    autosize(ws, [16, 55, 45, 35])

    # ─── 2. Masalah & Keputusan ───
    ws2 = wb.create_sheet("2. Masalah & Keputusan")
    ws2["A1"] = "Apa yang belum / bermasalah + keputusan"
    ws2["A1"].font = TITLE
    ws2.append([])
    ws2.append(["#", "Masalah (bahasa biasa)", "Dampak ke user", "Prioritas"])
    style_header(ws2, 4, 3)
    masalah = [
        (1, 'Saat kirim klaim terakhir, sistem bilang total "matching" tidak sama dengan total klaim', "Klaim stuck di draft; proses tidak selesai", "Paling penting"),
        (2, "User bisa lanjut simpan/submit meski data belum lengkap (kontrak tanpa aktivitas, approval tanpa centang)", "Risiko dokumen kosong / salah tersimpan", "Penting"),
        (3, "ABM sulit mencari klaim lewat kotak pencarian (harus buka lewat tautan ID)", "Ribet untuk verifikasi harian", "Penting"),
        (4, "Persentase fee tidak bisa diubah lewat layar pengaturan MAVEN", "Setting fee hanya lewat database / seed", "Sedang"),
    ]
    for i, row in enumerate(masalah):
        rr = 4 + i
        for c, val in enumerate(row, 1):
            ws2.cell(rr, c, val)
        ws2.cell(rr, 4).fill = RED if row[3] == "Paling penting" else (YELLOW if row[3] == "Penting" else BLUE)
    border_range(ws2, 3, 7, 1, 4)

    ws2.cell(9, 1, "Keputusan yang diminta").font = BOLD
    keputusan = [
        "1. Prioritas perbaikan #1: perbaiki validasi/pencocokan matching saat Submit klaim agar alur sampai akhir.",
        "2. Prioritas #2: perketat validasi sebelum dialog “Yakin simpan/submit?”.",
        "3. Prioritas #3: perbaiki pencarian klaim untuk ABM.",
    ]
    for i, t in enumerate(keputusan):
        ws2.cell(10 + i, 1, t)
        ws2.merge_cells(start_row=10 + i, start_column=1, end_row=10 + i, end_column=4)
        ws2.cell(10 + i, 1).alignment = WRAP
    ws2.cell(14, 1, "Catatan: beberapa “gagal” di angka tes otomatis bukan bug aplikasi (login role / tombol sengaja disembunyikan). Lihat sheet teknis.")
    ws2.cell(14, 1).font = Font(italic=True, size=10, color="666666")
    autosize(ws2, [5, 70, 40, 16])

    # ─── 3. Temuan Dev (ticket) ───
    ws3 = wb.create_sheet("3. Temuan Dev (Ticket)")
    ws3["A1"] = "Temuan prioritas — buat ticket perbaikan"
    ws3["A1"].font = TITLE
    ws3.append([])
    ws3.append(["ID", "Severity", "Temuan", "Dampak", "Rekomendasi"])
    style_header(ws3, 5, 3)
    tickets = [
        ("S5 / SV-CF-02", "Critical", "Alert: Total matching harus sama dengan total klaim! saat Submit", "Klaim tetap DRAFT setelah Ready=Y", "Pastikan alokasi Umbrand/SKU/CMA = total invoice; atau auto-match All Brand"),
        ("SV-MK-01 / SV-MK-02", "High", "Submit MKPP tanpa activity / Program Desc masih muncul Submit this data?", "Dokumen invalid bisa dikonfirmasi user", "Validasi wajib sebelum bootbox confirm"),
        ("SV-AP-02", "High", "Approval Save tanpa centang baris masih Save this data?", "No-op / risiko data kacau", "Block Save jika 0 Approve & 0 Reject"),
        ("FN-KL-02", "High", "LOV Find KlaimDF search by ID kosong (ABM)", "User & automation susah cari klaim", "Perbaiki filter LOV / search DocNo; workaround KlaimDF?ID="),
        ("CODE-02 / FN-MAP-13", "Med", "FeePercentage UI save di-comment", "Fee % tidak bisa diubah dari Mapping", "Uncomment + uji regresi"),
        ("FN-MK-19 / FN-PR-21 / FN-AP-09", "Low", "Stub: Update Attachment MKPP, ValidateFeeBalance PRM, Add Row Approval", "Fitur tidak tersedia", "Dokumentasikan by-design atau enable"),
        ("MAVEN ChooseRole", "Med (harness)", "Smoke Mapping/Monitoring goyah di ChooseRole", "False Fail Monitoring", "Stabilkan login + pilih role SUPER USER di runner"),
    ]
    sev_fill = {"Critical": RED, "High": ORANGE, "Med": YELLOW, "Low": BLUE, "Med (harness)": BLUE}
    for i, row in enumerate(tickets):
        rr = 4 + i
        for c, val in enumerate(row, 1):
            ws3.cell(rr, c, val)
        ws3.cell(rr, 2).fill = sev_fill.get(row[1], YELLOW)
        ws3.cell(rr, 2).font = BOLD
    border_range(ws3, 3, 10, 1, 5)
    autosize(ws3, [22, 14, 55, 35, 45])

    # ─── 4. Alur & Angka ───
    ws4 = wb.create_sheet("4. Alur & Angka")
    ws4["A1"] = "Alur teknis + angka hasil sesi"
    ws4["A1"].font = TITLE
    ws4["A3"] = "Alur end-to-end"
    ws4["A3"].font = BOLD
    ws4.append([])
    ws4.append(["Urutan", "Tahap", "Status", "Detail"])
    style_header(ws4, 4, 5)
    alur = [
        (1, "MAVEN fee API (CSA Offer 0.5% / Ads 0%)", "PASS", ""),
        (2, "MKPP create ≤100k (multi-fee / preset)", "PASS", ""),
        (3, "Patrick approve (single-tier)", "PASS", "Tanpa Hendrawan"),
        (4, "PRM Save klaim", "PASS", "ID 5330 / 26.10/KLAIM-CSA/036"),
        (5, "fahmi.permadi Ready=Y", "PASS", ""),
        (6, "Muslih Submit klaim", "FAIL", "S5 matching"),
    ]
    for i, row in enumerate(alur):
        rr = 6 + i
        for c, val in enumerate(row, 1):
            ws4.cell(rr, c, val)
        ws4.cell(rr, 3).fill = GREEN if row[2] == "PASS" else RED
        ws4.cell(rr, 3).font = BOLD
    border_range(ws4, 5, 11, 1, 4)

    ws4.cell(13, 1, "Roles uji").font = BOLD
    ws4.cell(14, 1, "Muslih → PATRICK.YOHANNES → PRM Subdist → fahmi.permadi → Muslih")

    ws4.cell(16, 1, "Angka hasil sesi 2026_10_18_10-02-29").font = BOLD
    ws4.cell(17, 1, "Jenis uji")
    ws4.cell(17, 2, "Dicek")
    ws4.cell(17, 3, "Pass")
    ws4.cell(17, 4, "Fail")
    ws4.cell(17, 5, "Stub / lain")
    style_header(ws4, 5, 17)
    ws4.cell(18, 1, "Track 2 — fungsi form")
    ws4.cell(18, 2, 40)
    ws4.cell(18, 3, 29)
    ws4.cell(18, 4, 7)
    ws4.cell(18, 5, "4 Stub")
    ws4.cell(19, 1, "Track 1 — skema varian")
    ws4.cell(19, 2, 14)
    ws4.cell(19, 3, 10)
    ws4.cell(19, 4, 4)
    ws4.cell(19, 5, 0)
    border_range(ws4, 17, 19, 1, 5)

    ws4.cell(21, 1, "Cara baca Fail: tidak semua Fail = bug produk. Ada Fail harness (ChooseRole), Fail by-design (Submit PRM hidden saat form kosong), dan Fail produk sejati (S5, validasi generik, LOV ABM).")
    ws4.merge_cells("A21:E22")
    ws4.cell(21, 1).alignment = WRAP
    autosize(ws4, [10, 45, 10, 35, 12])

    # ─── 5. Per Form ───
    ws5 = wb.create_sheet("5. Hasil per Form")
    ws5["A1"] = "Track 2 — kesimpulan per form"
    ws5["A1"].font = TITLE
    ws5.append([])
    ws5.append(["Form", "Kesimpulan", "Catatan"])
    style_header(ws5, 3, 3)
    forms = [
        ("MAVEN Mapping", "Belum andal di sesi ini", "Login ChooseRole mengganggu smoke; Fee % = Stub (CODE-02)"),
        ("MAVEN Monitoring", "Belum andal di sesi ini", "Sama — bukan conclusi bug produk sampai smoke ulang"),
        ("KDS MKPPDF", "Sehat", "New/Find/Save/Submit/LOV Group/Add Activity OK; Print missing di New; Update Att stub"),
        ("KDS ApprovalDoc", "Sehat", "Inbox + Approve + Save OK; Add Row stub"),
        ("PRM KlaimKDSDF", "Sehat untuk draft", "Save path OK; Submit hidden saat form kosong = wajar; ValidateFeeBalance stub"),
        ("KDS KlaimDF", "ABM OK, Submit CF blocked", "Ready/Reject/Memos OK; Find LOV ID gagal; Submit Muslih kena S5"),
    ]
    for i, row in enumerate(forms):
        rr = 4 + i
        for c, val in enumerate(row, 1):
            ws5.cell(rr, c, val)
        if "Sehat" in row[1] and "blocked" not in row[1].lower():
            ws5.cell(rr, 2).fill = GREEN
        elif "Belum" in row[1]:
            ws5.cell(rr, 2).fill = YELLOW
        else:
            ws5.cell(rr, 2).fill = ORANGE
    border_range(ws5, 3, 9, 1, 3)
    autosize(ws5, [22, 28, 75])

    # ─── 6. Skema Variants ───
    ws6 = wb.create_sheet("6. Skema Variants")
    ws6["A1"] = "Track 1 — hasil skema varian"
    ws6["A1"].font = TITLE
    ws6["A3"] = "Pass / sesuai desain"
    ws6["A3"].font = BOLD
    ws6.append([])
    ws6.append(["ID", "Isi", "Hasil", "Status"])
    style_header(ws6, 4, 5)
    passes = [
        ("SV-MV-01", "API GetActivitiesByGroup CSA", "200; Offer + Advertising ada", "Pass"),
        ("SV-MV-02", "groupName kosong", "data:null, bukan 500", "Pass"),
        ("SV-MK-09", "Cap >100.000 ditolak runner", "Pass", "Pass"),
        ("SV-CF-02", "Submit matching", "Alert matching muncul (bug dikonfirmasi)", "Pass*"),
        ("SV-MK-05 / AP-01 / PR-01 / AB-01", "Ref happy path", "Doc 0054–0057, klaim 5330 Ready=Y", "Pass"),
        ("SV-AB-02 / AB-06", "LOV ID kosong; New/Submit ABM hidden", "Temuan UX + role gate OK", "Pass"),
    ]
    for i, row in enumerate(passes):
        rr = 6 + i
        for c, val in enumerate(row, 1):
            ws6.cell(rr, c, val)
        ws6.cell(rr, 4).fill = GREEN
    border_range(ws6, 5, 11, 1, 4)
    ws6.cell(12, 1, "* SV-CF-02 Pass = skenario negatif berhasil mendeteksi bug matching (bukan berarti Submit klaim lulus).")
    ws6.cell(12, 1).font = Font(italic=True, size=9, color="666666")

    ws6.cell(14, 1, "Fail produk / gap").font = BOLD
    ws6.cell(15, 1, "ID")
    ws6.cell(15, 2, "Isi")
    ws6.cell(15, 3, "Aktual")
    ws6.cell(15, 4, "Status")
    style_header(ws6, 4, 15)
    fails = [
        ("SV-MK-01", "Submit tanpa activity", "Konfirmasi generik, bukan validasi activity", "Fail"),
        ("SV-MK-02", "Submit tanpa Program Desc", "Sama", "Fail"),
        ("SV-AP-02", "Approval tanpa centang", "Konfirmasi generik", "Fail"),
        ("SV-PR-03", "Faktur 16 digit", "Belum sampai assert digit; kena Subdist belum dipilih (setup)", "Fail"),
    ]
    for i, row in enumerate(fails):
        rr = 16 + i
        for c, val in enumerate(row, 1):
            ws6.cell(rr, c, val)
        ws6.cell(rr, 4).fill = RED
    border_range(ws6, 15, 19, 1, 4)
    autosize(ws6, [28, 40, 50, 10])

    # ─── 7. Data Uji & Lampiran ───
    ws7 = wb.create_sheet("7. Data Uji & Lampiran")
    ws7["A1"] = "Data uji DEV + cara baca lampiran"
    ws7["A1"].font = TITLE
    ws7.append([])
    ws7.append(["Dokumen", "Nilai", "Kegunaan"])
    style_header(ws7, 3, 3)
    data = [
        ("26.10/MKPP-MIM/0054.000", "Rp 100.000 mixed fee", "Happy path utama"),
        ("0055 / 0056 / 0057", "10k / 30k / 40k", "Batch preset"),
        ("26.10/KLAIM-CSA/036 (ID 5330)", "Ready=Y", "Repro S5 + fungsi KlaimDF"),
    ]
    for i, row in enumerate(data):
        rr = 4 + i
        for c, val in enumerate(row, 1):
            ws7.cell(rr, c, val)
    border_range(ws7, 3, 6, 1, 3)

    ws7.cell(8, 1, "Glosarium singkat").font = BOLD
    ws7.cell(9, 1, "Istilah")
    ws7.cell(9, 2, "Arti")
    style_header(ws7, 2, 9)
    gloss = [
        ("MKPP", "Kontrak / memo Development Fund di KICAO KDS"),
        ("Klaim", "Pengajuan klaim Subdist (KDS atau PRM)"),
        ("ABM", "Role verifikasi lapangan (uji: fahmi.permadi)"),
        ("Matching", "Alokasi rinci yang jumlahnya harus = total klaim"),
        ("Stub", "Fitur UI/API sengaja nonaktif di kode"),
        ("Track 1", "Variasi skenario (positif / negatif / edge)"),
        ("Track 2", "Cek tombol & fungsi per form (smoke)"),
    ]
    for i, row in enumerate(gloss):
        rr = 10 + i
        ws7.cell(rr, 1, row[0])
        ws7.cell(rr, 2, row[1])
    border_range(ws7, 9, 16, 1, 2)

    ws7.cell(18, 1, "File terkait").font = BOLD
    ws7.cell(19, 1, "Laporan ringkas (MD)")
    ws7.cell(19, 2, str(REPORTS / "LAPORAN_RINGKAS.md"))
    ws7.cell(20, 1, "Laporan lengkap (MD)")
    ws7.cell(20, 2, str(REPORTS / "LAPORAN_DUAL_TRACK_FINAL.md"))
    ws7.cell(21, 1, "Excel detail run otomatis")
    ws7.cell(21, 2, str(REPORTS / "DF_DualTrack_Report_2026_10_18_10-02-29.xlsx"))
    ws7.cell(22, 1, "Folder evidence / screenshot")
    ws7.cell(22, 2, str(EVIDENCE))
    if EVIDENCE.exists():
        c = ws7.cell(22, 2)
        c.hyperlink = str(EVIDENCE)
        c.font = Font(color="0563C1", underline="single")

    ws7.cell(24, 1, "Cara reproduksi").font = BOLD
    ws7.cell(25, 1, "cd Project/DevelopmentFund/Testing/runners")
    ws7.cell(26, 1, "py -3 run_dual_track.py")
    autosize(ws7, [32, 70, 30])

    wb.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
