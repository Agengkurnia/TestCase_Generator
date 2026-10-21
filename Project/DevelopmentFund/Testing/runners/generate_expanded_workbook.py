"""Generate Excel workbook from cases_expanded.py (standard column set)."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from _pack import REPORTS
from cases_expanded import CASES, MKPP_PRESETS, BUDGET_CAP

HERE = REPORTS
STAMP = datetime.now().strftime("%Y_%m_%d_%H-%M-%S")
QA = "Ageng"
TODAY = datetime.now().strftime("%d/%m/%Y")

HEADERS = [
    "No.",
    "Scenario ID",
    "Modul / Area",
    "Field Skenario",
    "Tipe",
    "Deskripsi",
    "Prasyarat",
    "Langkah",
    "Expected",
    "Priority",
    "Status Uji",
    "Hasil Aktual",
    "Bukti Pengujian",
    "QA",
    "Tanggal Uji",
]


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Expanded Cases"
    head_fill = PatternFill("solid", fgColor="1E7D37")
    head_font = Font(color="FFFFFF", bold=True, size=10)
    thin = Border(
        left=Side(style="thin", color="E5E7EB"),
        right=Side(style="thin", color="E5E7EB"),
        top=Side(style="thin", color="E5E7EB"),
        bottom=Side(style="thin", color="E5E7EB"),
    )
    ws.append(HEADERS)
    for col, _ in enumerate(HEADERS, 1):
        c = ws.cell(1, col)
        c.fill = head_fill
        c.font = head_font
        c.alignment = Alignment(wrap_text=True, vertical="center")

    for i, tc in enumerate(CASES, 1):
        ws.append(
            [
                i,
                tc["id"],
                tc["modul"],
                tc["field"],
                tc["tipe"],
                tc["desc"],
                tc["prereq"],
                tc["langkah"],
                tc["expected"],
                tc["priority"],
                "Not Run",
                "",
                "",
                QA,
                TODAY,
            ]
        )
        for col in range(1, len(HEADERS) + 1):
            cell = ws.cell(i + 1, col)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = thin

    widths = [5, 22, 14, 28, 10, 40, 28, 40, 32, 10, 12, 24, 18, 10, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:O{len(CASES) + 1}"

    ws2 = wb.create_sheet("MKPP Presets (cap)")
    ws2.append(["Preset ID", "Title", "Group", "Activities", "Total", "SKP/act", "Under Cap?"])
    for p in MKPP_PRESETS:
        total = sum(int(a["amount"]) for a in p["activities"])
        acts = " | ".join(f"{a['outlet']}:{a['act_code']}={a['amount']}" for a in p["activities"])
        ws2.append([p["id"], p["title"], p["group"], acts, total, p["skp_per_act"], "YES" if total <= BUDGET_CAP else "NO"])
    ws2["A1"].font = Font(bold=True)

    ws3 = wb.create_sheet("Notes")
    ws3["A1"] = "Budget rule"
    ws3["B1"] = f"Setiap MKPP sum(activity.amount) <= {BUDGET_CAP:,}".replace(",", ".")
    ws3["A2"] = "Roles"
    ws3["B2"] = "MKPP approve=PATRICK.YOHANNES only; ABM=fahmi.permadi; Creator/Submit klaim=Muslih"
    ws3["A3"] = "Generated"
    ws3["B3"] = STAMP
    ws3["A4"] = "Case count"
    ws3["B4"] = len(CASES)

    out = HERE / f"DF_SIT_Expanded_Cases_{STAMP}.xlsx"
    wb.save(out)
    print(f"Wrote {out} ({len(CASES)} cases, {len(MKPP_PRESETS)} presets, cap={BUDGET_CAP})")


if __name__ == "__main__":
    main()
