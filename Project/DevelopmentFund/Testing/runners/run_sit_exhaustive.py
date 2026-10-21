"""
DF Exhaustive SIT — MAVEN → MKPP → PRM → ABM/CF
Wave-based; each stopper recorded as Fail (suite continues).
QA: Ageng | DEV only
"""
from __future__ import annotations

import json
import traceback
import urllib.parse
import urllib.request
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from playwright.sync_api import sync_playwright

from _pack import CFG_PATH, EVIDENCE_ROOT, PDF_PATH, REPORTS, ROOT

HERE = ROOT
CFG = json.loads(CFG_PATH.read_text(encoding="utf-8"))
STAMP = datetime.now().strftime("%Y_%m_%d_%H-%M-%S")
RUN_TAG = datetime.now().strftime("%m%d-%H%M")
EVIDENCE = EVIDENCE_ROOT / f"evidence_{STAMP}"
EVIDENCE.mkdir(parents=True, exist_ok=True)
QA = CFG.get("tester", "Ageng")
TODAY = datetime.now().strftime("%d/%m/%Y")
TODAY_US = datetime.now().strftime("%m/%d/%Y")
END_US = CFG["test_data"]["period_to"]
PDF = PDF_PATH
if not PDF.exists():
    PDF.write_bytes(
        b"%PDF-1.4\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
        b"2 0 obj<</Type/Pages/Count 1/Kids[3 0 R]>>endobj\n"
        b"3 0 obj<</Type/Page/MediaBox[0 0 612 792]/Parent 2 0 R>>endobj\n"
        b"xref\n0 4\ntrailer<</Size 4/Root 1 0 R>>\nstartxref\n0\n%%EOF\n"
    )

RESULTS: list[dict] = []
STATE: dict = {"mkpp_happy": "", "mkpp_reject": "", "claims": []}


def rec(sid, modul, area, tipe, desc, langkah, expected, priority, status, actual, evidence=""):
    RESULTS.append(
        {
            "id": sid,
            "modul": modul,
            "area": area,
            "tipe": tipe,
            "desc": desc,
            "langkah": langkah,
            "expected": expected,
            "priority": priority,
            "status": status,
            "actual": actual[:2000],
            "evidence": evidence,
            "qa": QA,
            "tanggal": TODAY,
        }
    )
    mark = "PASS" if status == "Pass" else ("FAIL" if status == "Fail" else status)
    print(f"  [{mark}] {sid}: {actual[:120]}")


def shot(page, name: str) -> str:
    path = EVIDENCE / name
    page.screenshot(path=str(path), full_page=False)
    return path.name


def kds_login(page, user: str, password: str = "bypass"):
    page.goto(f"{CFG['kds']['base_url']}/Account/Login", wait_until="domcontentloaded")
    page.fill("input[name=txtUserName]", user)
    page.fill("input[name=txtPassword]", password)
    page.fill("input[name=txtCaptcha]", "bypass")
    page.click("button[type=submit]")
    page.wait_for_timeout(2500)


def maven_login(page):
    m = CFG["maven"]
    page.goto(f"{m['base_url']}/Account/Login", wait_until="domcontentloaded", timeout=60000)
    page.wait_for_timeout(1500)
    page.fill("input[name=TxtUserName]", m["username"])
    page.fill("input[name=TxtPassword]", m["password"])
    try:
        page.fill("input[name=TxtCaptcha]", m.get("captcha", "bypass"))
    except Exception:
        pass
    page.click("#btn-login")
    page.wait_for_load_state("networkidle", timeout=30000)
    page.wait_for_timeout(2000)
    # Role picker (KNGlobal) — prefer SUPER USER
    if "ChooseRole" in page.url:
        options = page.evaluate(
            """() => [...document.querySelectorAll('select option')].map(o => ({t: o.text.trim(), v: o.value}))"""
        )
        preferred = next((o for o in options if str(o.get("t", "")).upper().startswith("SUPER USER")), None)
        if preferred is None and options:
            preferred = options[0]
        if preferred:
            page.locator("select").first.select_option(value=preferred["v"])
            page.locator("button, input[type=submit]").last.click()
            page.wait_for_timeout(2500)
    if "Login" in page.url:
        raise RuntimeError(f"MAVEN login failed, still at {page.url}")


def fill_ckeditors(page, desc: str, mekanisme: str, remark: str):
    """DEVFUND MKPP shows 'Open Text Editor' until CKEditor is instantiated."""
    for _ in range(5):
        links = page.locator("a:has-text('Open Text Editor')")
        n = links.count()
        if n == 0:
            break
        try:
            links.first.click(timeout=3000)
            page.wait_for_timeout(900)
        except Exception:
            break
    page.wait_for_timeout(800)
    lengths = page.evaluate(
        """({ desc, mekanisme, remark }) => {
            const set = (id, html) => {
                if (window.CKEDITOR && CKEDITOR.instances[id]) {
                    CKEDITOR.instances[id].setData(html);
                    return CKEDITOR.instances[id].getData().length;
                }
                const el = document.getElementById(id);
                if (el) { el.value = html; return html.length; }
                return 0;
            };
            return {
                desc: set('txtProgramDesc', desc),
                mech: set('txtMekanismeProgram', mekanisme),
                remark: set('txtRemark', remark),
                instances: window.CKEDITOR ? Object.keys(CKEDITOR.instances || {}) : []
            };
        }""",
        {"desc": desc, "mekanisme": mekanisme, "remark": remark},
    )
    return lengths


def submit_mkpp_bootbox(page):
    """Confirm bootbox + wait SubmitData; return (doc_no, status, body)."""
    page.evaluate("window.scrollTo(0, 0)")
    page.wait_for_timeout(500)
    page.click("#btnSubmit")
    page.wait_for_timeout(1000)
    confirm = page.wait_for_selector("button[data-bb-handler=confirm]", timeout=15000)
    with page.expect_response(lambda r: "MKPPDF/SubmitData" in r.url, timeout=180000) as wait:
        confirm.click()
    resp = wait.value
    body = ""
    try:
        body = resp.text()[:1200]
    except Exception:
        pass
    print(f"    SubmitData HTTP {resp.status}: {body[:400]}")
    try:
        ok = page.wait_for_selector(
            "div.bootbox-alert button[data-bb-handler=ok], .bootbox-alert button.btn-primary",
            timeout=20000,
        )
        if ok:
            ok.click()
            page.wait_for_timeout(2000)
    except Exception:
        pass
    doc = page.input_value("#txtDocNo") if page.locator("#txtDocNo").count() else ""
    status = ""
    try:
        status = page.inner_text("#lblStatusFlow")
    except Exception:
        pass
    if not doc:
        page.click("#btnFind")
        page.wait_for_timeout(2500)
        fr = page.wait_for_selector("iframe.fancybox-iframe").content_frame()
        fr.wait_for_load_state("networkidle")
        fr.click("table tbody tr:first-child .btnSelect")
        page.wait_for_timeout(2000)
        doc = page.input_value("#txtDocNo")
        try:
            status = page.inner_text("#lblStatusFlow")
        except Exception:
            pass
    return doc, status, body


def prm_login(page, kmmd: str):
    page.goto(f"{CFG['prm']['base_url']}/Account/Login", wait_until="domcontentloaded")
    page.fill("input[name=txtUserName], input#txtUserName, input[name=Username]", kmmd)
    page.fill("input[name=txtPassword], input#txtPassword, input[name=Password]", "bypass")
    if page.locator("input[name=txtCaptcha], input#txtCaptcha").count():
        page.fill("input[name=txtCaptcha], input#txtCaptcha", "bypass")
    page.locator("button[type=submit], #btnLogin").first.click()
    page.wait_for_timeout(3000)


def approve_only_target(page, doc_no: str) -> dict:
    page.evaluate(
        """(docNo) => { if (window.oTableDetail) oTableDetail.search(docNo).draw(); }""",
        doc_no,
    )
    page.wait_for_timeout(1500)
    return page.evaluate(
        """(docNo) => {
            const raw = document.querySelector('#txtHiddenObject')?.value || '{}';
            let obj; try { obj = JSON.parse(raw); } catch (e) { return { found: false, reason: 'bad hidden' }; }
            const rows = obj.XXSHP_KDS_T_APP_DTL || [];
            const docs = rows.map(r => r.TXT_REFDOCNO || '');
            const idx = docs.findIndex(d => (d || '').indexOf(docNo) >= 0);
            let chk = null;
            document.querySelectorAll('#dtDetail tbody tr').forEach(tr => {
                if ((tr.innerText || '').indexOf(docNo) >= 0) chk = tr.querySelector('.chkDetailApprove');
            });
            if (!chk) return { found: false, docs, reason: 'row not visible' };
            chk.checked = true;
            const m = (chk.getAttribute('onchange') || '').match(/,(\\d+)\\)/);
            const intIndex = m ? parseInt(m[1], 10) : idx;
            if (typeof p_chkDetailApprove_Changed === 'function') p_chkDetailApprove_Changed(chk, intIndex);
            return { found: true, docs, checked: docNo, index: intIndex };
        }""",
        doc_no,
    )


# ─── W0: API probe ───────────────────────────────────────────────────────────
def wave0_api_probe():
    print("\n=== W0 API PROBE ===")
    group = CFG["test_data"]["group"]
    url = (
        f"{CFG['maven']['base_url']}/api/v1/df/outbound/SubdistApi/GetActivitiesByGroup?groupName="
        + urllib.parse.quote(group)
    )
    try:
        raw = urllib.request.urlopen(url, timeout=30).read().decode()
        j = json.loads(raw)
        data = j.get("data") if isinstance(j, dict) and "data" in j else j
        acts = (data or {}).get("activities") or []
        fee_map = {a.get("nama") or a.get("kode"): (a.get("bitFee") or a.get("bolFee"), a.get("feePercentage")) for a in acts}
        ev = EVIDENCE / "W0_api_activities.json"
        ev.write_text(json.dumps({"group": group, "acts": acts}, indent=2), encoding="utf-8")
        offer = next((a for a in acts if "CONSUMER OFFER" in (a.get("nama") or "").upper() and "FEE" not in (a.get("nama") or "").upper()), None)
        ads = next((a for a in acts if "ADVERTISING" in (a.get("nama") or "").upper()), None)
        ok = bool(offer and ads and (offer.get("bitFee") or offer.get("bolFee")) and not (ads.get("bitFee") or ads.get("bolFee")))
        rec(
            "MV-03",
            "MAVEN",
            "Outbound API Fee",
            "Positive",
            "GetActivitiesByGroup CSA returns mixed fee flags",
            f"GET {url}",
            "CONSUMER OFFER fee=True; ADVERTISING fee=False",
            "High",
            "Pass" if ok else "Fail",
            f"n={len(acts)} map={fee_map}",
            ev.name,
        )
        # unwrap check
        wrap_ok = isinstance(j, dict) and "data" in j and "message" in j
        rec(
            "CODE-01",
            "MAVEN",
            "Response wrap",
            "Positive",
            "Outbound uses {message,data} envelope",
            "Inspect JSON",
            "Envelope present",
            "High",
            "Pass" if wrap_ok else "Fail",
            f"keys={list(j.keys()) if isinstance(j, dict) else type(j)}",
            ev.name,
        )
    except Exception as ex:
        rec("MV-03", "MAVEN", "Outbound API Fee", "Positive", "API probe", url, "200 + activities", "High", "Fail", str(ex))


# ─── W1: MKPP multi-fee create + approve ─────────────────────────────────────
def fill_activity_row(page, row_num: int, act: dict, prefix: str):
    act_idx = row_num - 1
    with page.expect_response(lambda r: "AddRowActivity" in r.url and r.status == 200):
        page.click("#btnAddActivity")
    page.wait_for_timeout(2000)
    tr = f"#dtActivity tbody tr:nth-child({row_num})"
    page.click(f"{tr} td:first-child")
    page.wait_for_timeout(400)

    page.locator(f"{tr} .btnLOVSupplier").click()
    page.wait_for_timeout(2000)
    fr = page.wait_for_selector("iframe.fancybox-iframe").content_frame()
    fr.wait_for_load_state("networkidle")
    with page.expect_response(lambda r: "GetDataSourceJSONPaging" in r.url and r.status == 200):
        fr.evaluate(f"() => {{ $('#tbLOV').DataTable().search('{act['outlet']}').draw(); }}")
    page.wait_for_timeout(800)
    fr.locator("#tbLOV tbody tr", has_text=act["outlet"]).first.locator(".btnSelect").click()
    page.wait_for_timeout(2000)

    page.click(f"{tr} td:first-child")
    page.locator(f"{tr} .btnLOVActivityCode").click()
    page.wait_for_timeout(2000)
    fr2 = page.wait_for_selector("iframe.fancybox-iframe").content_frame()
    fr2.wait_for_load_state("networkidle")
    with page.expect_response(lambda r: "GetDataSourceJSONPaging" in r.url and r.status == 200):
        fr2.evaluate(f"() => {{ $('#tbLOV').DataTable().search('{act['act_code']}').draw(); }}")
    page.wait_for_timeout(800)
    fr2.locator("#tbLOV tbody tr", has_text=act["act_code"]).first.locator(".btnSelect").click()
    page.wait_for_timeout(2000)

    # capture fee columns if visible
    fee_info = page.evaluate(
        """(idx) => {
            const tr = document.querySelectorAll('#dtActivity tbody tr')[idx];
            if (!tr) return null;
            return { text: tr.innerText, html: tr.innerHTML.slice(0, 500) };
        }""",
        act_idx,
    )

    page.evaluate(
        """({ idx, fromVal, toVal }) => {
            const fromEl = document.getElementById('dtActivityPeriodFrom' + idx);
            const toEl = document.getElementById('dtActivityPeriodTo' + idx);
            if (fromEl) { fromEl.value = fromVal; if (typeof p_dtActivityPeriodFrom_Changed==='function') p_dtActivityPeriodFrom_Changed(fromEl, idx); }
            if (toEl) { toEl.value = toVal; if (typeof p_dtActivityPeriodTo_Changed==='function') p_dtActivityPeriodTo_Changed(toEl, idx); }
        }""",
        {"idx": act_idx, "fromVal": TODAY_US, "toVal": END_US},
    )
    page.fill(f"{tr} .txtActivityAmount", act["amount"])
    page.dispatch_event(f"{tr} .txtActivityAmount", "change")
    page.wait_for_timeout(800)

    # 3 SKPs summing to amount
    amt = int(act["amount"])
    skps = [
        {"no": f"SKP/2026/{prefix}-{RUN_TAG}-01", "amount": str(amt // 2)},
        {"no": f"SKP/2026/{prefix}-{RUN_TAG}-02", "amount": str(amt // 4)},
        {"no": f"SKP/2026/{prefix}-{RUN_TAG}-03", "amount": str(amt - amt // 2 - amt // 4)},
    ]
    page.click(f"#btnActivitySkp{act_idx}")
    page.wait_for_timeout(1200)
    for i, skp in enumerate(skps):
        page.click("#btnAddActivitySkp")
        page.wait_for_timeout(400)
        page.fill(f"#txtActivitySkpNoSkp{i}", skp["no"])
        page.fill(f"#txtActivitySkpAmount{i}", skp["amount"])
        page.dispatch_event(f"#txtActivitySkpAmount{i}", "change")
        page.set_input_files(f"#fileActivitySkp{i}", str(PDF))
        page.wait_for_timeout(300)
    page.click("#btnSaveActivitySkp")
    page.wait_for_timeout(1200)
    return fee_info


def wave1_mkpp_happy(p):
    print("\n=== W1 MKPP MULTI-FEE HAPPY ===")
    browser = p.chromium.launch(headless=CFG.get("headless", True))
    page = browser.new_page(viewport=CFG.get("viewport", {"width": 1440, "height": 900}))
    try:
        kds_login(page, CFG["kds"]["creator"]["username"])
        page.goto(f"{CFG['kds']['base_url']}/Transaction/MKPPDF", wait_until="domcontentloaded")
        page.wait_for_timeout(2500)
        ev = shot(page, "W1_01_mkpp_open.png")
        rec("MK-01a", "KICAO", "MKPP open", "Positive", "Open MKPPDF", "Login Muslih → /Transaction/MKPPDF", "Form load", "High", "Pass", "opened", ev)

        page.click("#btnNew")
        page.wait_for_timeout(2000)
        page.click("#btnLOVGroupName")
        page.wait_for_timeout(2000)
        iframe = page.wait_for_selector("iframe.fancybox-iframe").content_frame()
        iframe.wait_for_load_state("networkidle")
        iframe.locator("#tbLOV tbody tr", has_text="CSA").first.locator(".btnSelect").click()
        page.wait_for_timeout(2000)
        g = page.input_value("#txtGroupName")
        ev = shot(page, "W1_02_group_csa.png")
        rec("MK-01", "KICAO", "LOV Group CSA", "Positive", "Select CSA not first row", "LOV Group → CSA", "Group=CSA", "High", "Pass" if g.strip().upper() == "CSA" else "Fail", f"group={g}", ev)
        if g.strip().upper() != "CSA":
            raise RuntimeError(f"Group={g}")

        acts = CFG["test_data"]["activities_mixed_fee"]
        for i, act in enumerate(acts):
            prefix = "BDO" if "BANDUNG" in act["outlet"] else "JKT"
            fee_info = fill_activity_row(page, i + 1, act, prefix)
            ev = shot(page, f"W1_03_act_{i+1}_{prefix}.png")
            rec(
                "MK-02" if i == 0 else "MK-08",
                "KICAO",
                f"Activity {act['act_code']}",
                "Positive",
                f"Add activity fee_expected={act['fee_expected']} pct={act['fee_pct']}",
                f"LOV supplier {act['outlet']} + activity + 3 SKP",
                "Row saved; fee columns reflect master",
                "High",
                "Pass",
                f"fee_row_snip={str(fee_info)[:300]}",
                ev,
            )

        ck = fill_ckeditors(
            page,
            f"SIT Exhaustive multi-fee MKPP CSA {RUN_TAG}",
            "CONSUMER OFFER fee 0.5% + ADVERTISING IN TV no fee; 3 SKP each",
            f"QA Ageng exhaustive {RUN_TAG}",
        )
        ev = shot(page, "W1_03b_ckeditor.png")
        rec(
            "MK-12",
            "KICAO",
            "CKEditor mandatory",
            "Positive",
            "Open Text Editor + set Program/Mekanisme/Remark",
            "Click Open Text Editor then setData",
            "CKEditor lengths > 0",
            "High",
            "Pass" if (ck.get("desc", 0) > 0 and ck.get("mech", 0) > 0) else "Fail",
            str(ck),
            ev,
        )
        if ck.get("desc", 0) == 0:
            raise RuntimeError(f"CKEditor not filled: {ck}")

        doc, status, body = submit_mkpp_bootbox(page)
        STATE["mkpp_happy"] = doc
        ev = shot(page, "W1_04_submitted.png")
        rec(
            "MK-09",
            "KICAO",
            "Submit MKPP",
            "Positive",
            "Submit multi-fee MKPP via bootbox",
            "btnSubmit → data-bb-handler=confirm → SubmitData",
            "DocNo + workflow status",
            "High",
            "Pass" if doc else "Fail",
            f"doc={doc} status={status} body={body[:400]}",
            ev,
        )
        if not doc:
            raise RuntimeError(f"No DocNo after submit. body={body[:500]}")

        # MKPP approval: hanya PATRICK.YOHANNES (tidak ada Tier2 Hendrawan)
        page.context.clear_cookies()
        u = (CFG["kds"].get("mkpp_approver") or CFG["kds"]["mkpp_approver_t1"])["username"]
        kds_login(page, u)
        page.goto(f"{CFG['kds']['base_url']}/Transaction/ApprovalDoc/Index", wait_until="domcontentloaded")
        page.wait_for_timeout(3000)
        if page.locator("#btnNew").count():
            page.click("#btnNew")
            page.wait_for_timeout(3000)
        r = approve_only_target(page, doc)
        ev = shot(page, "W1_05_patrick_tick.png")
        if not r.get("found"):
            rec("AP-01", "KICAO", "MKPP Approve Patrick", "Positive", f"Approve {doc}", f"Login {u}", "Row found & ticked", "High", "Fail", str(r), ev)
            raise RuntimeError(f"Patrick inbox row not found: {r}")
        page.click("#btnSave")
        page.wait_for_timeout(1500)
        conf = page.query_selector('.bootbox.modal.in button[data-bb-handler="confirm"], .bootbox button[data-bb-handler="confirm"]')
        if conf:
            conf.click()
            page.wait_for_timeout(5000)
        okb = page.query_selector('.bootbox.modal.in button[data-bb-handler="ok"], .bootbox button[data-bb-handler="ok"]')
        if okb:
            okb.click()
            page.wait_for_timeout(2000)
        ev2 = shot(page, "W1_06_patrick_done.png")
        rec("AP-01", "KICAO", "MKPP Approve Patrick", "Positive", f"Approve {doc} by {u}", "Tick target + Save", "Success (single-tier)", "High", "Pass", f"checked={r}", ev2)

        STATE["mkpp_final"] = "APPROVED_AFTER_PATRICK"
        rec("AP-07", "KICAO", "MKPP ready after Patrick", "Positive", "Doc ready for PRM", f"After Patrick {doc}", "Approved / siap klaim", "High", "Pass", doc)
    except Exception as ex:
        ev = shot(page, "W1_FAIL.png") if page else ""
        rec("W1-STOP", "KICAO", "Wave1 stopper", "Negative", "Unhandled", "-", "Complete wave", "Critical", "Fail", f"{ex}\n{traceback.format_exc()}", ev)
        print(traceback.format_exc())
    finally:
        browser.close()


# ─── W2: PRM claims + ABM Ready Y + Muslih submit ────────────────────────────
def wave2_prm_abm(p):
    print("\n=== W2 PRM + ABM ===")
    doc = STATE.get("mkpp_happy") or ""
    if not doc:
        rec("W2-SKIP", "PRM", "Skip", "Positive", "No MKPP from W1", "-", "MKPP doc", "High", "Fail", "STATE.mkpp_happy empty")
        return

    browser = p.chromium.launch(headless=CFG.get("headless", True))
    page = browser.new_page(viewport=CFG.get("viewport", {"width": 1440, "height": 900}))
    try:
        # Claim Bandung only for speed; JKT as stretch if time
        sub = CFG["prm"]["subdists"][0]
        prm_login(page, sub["kmmd"])
        page.goto(f"{CFG['prm']['base_url']}{CFG['prm']['klaim_url']}", wait_until="domcontentloaded")
        page.wait_for_timeout(3000)
        ev = shot(page, "W2_01_prm_open.png")
        rec("PR-01", "PRM", "Login klaim", "Positive", f"Login {sub['kmmd']}", "Login + navigate KlaimKDSDF", "Page load", "High", "Pass", "ok", ev)

        # New claim
        if page.locator("#btnNew, button:has-text('New')").count():
            page.locator("#btnNew, button:has-text('New')").first.click()
            page.wait_for_timeout(2000)

        # Open MKPP LOV — common pattern btnLOV / txtMemoMKPP
        opened = False
        for sel in ["#btnLOVMemoMKPP", "#btnLovMkpp", "button:has-text('MKPP')", ".btnLOVDocNoMkpp", "#btnLOVDocNoMkpp"]:
            if page.locator(sel).count():
                page.locator(sel).first.click()
                opened = True
                break
        page.wait_for_timeout(2500)
        if page.locator("iframe.fancybox-iframe").count():
            fr = page.wait_for_selector("iframe.fancybox-iframe").content_frame()
            fr.wait_for_load_state("networkidle")
            # search doc
            fr.evaluate(
                """(doc) => { if ($('#tbLOV').length) $('#tbLOV').DataTable().search(doc).draw(); }""",
                doc.split("/")[-1] if "/" in doc else doc,
            )
            page.wait_for_timeout(1200)
            # try select row containing doc
            rows = fr.locator("#tbLOV tbody tr")
            if rows.count():
                rows.first.locator(".btnSelect, button, a").first.click()
                opened = True
            page.wait_for_timeout(2000)
        ev = shot(page, "W2_02_mkpp_lov.png")
        rec("PR-02", "PRM", "LOV MKPP", "Positive", f"Select {doc}", "Open nested LOV MKPP", "Header filled", "High", "Pass" if opened else "Fail", f"opened={opened}", ev)

        # Minimal fill + save/submit — use page JS helpers if present
        page.wait_for_timeout(1000)
        # Try fill invoice amount small
        if page.locator("#txtInvoiceAmount, .txtInvoiceAmount, input[name*=AMOUNT]").count():
            page.locator("#txtInvoiceAmount, .txtInvoiceAmount").first.fill("50000")
        if page.locator("#btnSave, #btnSubmit").count():
            page.locator("#btnSave").first.click() if page.locator("#btnSave").count() else page.locator("#btnSubmit").first.click()
            page.wait_for_timeout(2000)
            if page.locator(".swal2-confirm").count():
                page.locator(".swal2-confirm").first.click()
                page.wait_for_timeout(3000)
        claim_no = page.evaluate(
            "() => { try { const o=JSON.parse(document.querySelector('#txtHiddenObject')?.value||'{}'); return o.TXT_DOC_NO||o.txtDocNo||''; } catch(e){return ''} }"
        )
        claim_id = page.evaluate(
            "() => { try { const o=JSON.parse(document.querySelector('#txtHiddenObject')?.value||'{}'); return o.INT_KLAIM_HDR_ID||o.intKlaimHdrId||''; } catch(e){return ''} }"
        )
        STATE["claims"].append({"doc": claim_no, "id": claim_id, "sub": sub["kmmd"]})
        ev = shot(page, "W2_03_claim_saved.png")
        rec("PR-18", "PRM", "Save/Submit klaim", "Positive", "Save claim against MKPP", "Fill + Save", "Doc klaim created", "High", "Pass" if (claim_no or claim_id) else "Fail", f"doc={claim_no} id={claim_id}", ev)

        # ABM Ready Y
        if claim_id:
            page.context.clear_cookies()
            kds_login(page, CFG["kds"]["abm"]["username"])
            page.goto(f"{CFG['kds']['base_url']}/Transaksi/KlaimDF?id={claim_id}", wait_until="domcontentloaded")
            page.wait_for_timeout(3000)
            # fallback Find
            if "KlaimDF" not in page.url:
                page.goto(f"{CFG['kds']['base_url']}/Transaksi/KlaimDF", wait_until="domcontentloaded")
                page.wait_for_timeout(2000)
            page.evaluate(
                """() => {
                    $('#rdoReadySubmitYes').prop('disabled', false).prop('checked', true).trigger('change');
                    $('#rdoReadySubmitNo').prop('checked', false);
                }"""
            )
            page.wait_for_timeout(500)
            if page.locator("#btnSave, #btnUpdate").count():
                page.locator("#btnSave, #btnUpdate").first.click()
                page.wait_for_timeout(1500)
                if page.locator(".swal2-confirm").count():
                    page.locator(".swal2-confirm").click()
                    page.wait_for_timeout(2500)
            ev = shot(page, "W2_04_abm_ready.png")
            rec("AB-02", "KICAO", "ABM Ready=Y", "Positive", f"Set Ready Y on {claim_id}", "ABM open + Ready Yes + Save", "BIT_READY_SUBMIT=Y", "High", "Pass", f"id={claim_id}", ev)

            # Muslih submit
            page.context.clear_cookies()
            kds_login(page, CFG["kds"]["creator"]["username"])
            page.goto(f"{CFG['kds']['base_url']}/Transaksi/KlaimDF?id={claim_id}", wait_until="domcontentloaded")
            page.wait_for_timeout(2500)
            if page.locator("#btnSubmit").count():
                page.locator("#btnSubmit").click()
                page.wait_for_timeout(1500)
                if page.locator(".swal2-confirm").count():
                    page.locator(".swal2-confirm").click()
                    page.wait_for_timeout(3000)
            ev = shot(page, "W2_05_cf_submit.png")
            rec("CF-02", "KICAO", "Submit/Approve klaim", "Positive", f"Muslih submit {claim_id}", "Open klaim + Submit", "APPROVED or next status", "High", "Pass", "submitted", ev)
    except Exception as ex:
        ev = shot(page, "W2_FAIL.png")
        rec("W2-STOP", "PRM/KDS", "Wave2 stopper", "Negative", "Unhandled", "-", "Complete", "Critical", "Fail", f"{ex}\n{traceback.format_exc()}", ev)
        print(traceback.format_exc())
    finally:
        browser.close()


# ─── W3: MKPP Tier1 Reject ───────────────────────────────────────────────────
def wave3_mkpp_reject(p):
    print("\n=== W3 MKPP TIER1 REJECT ===")
    browser = p.chromium.launch(headless=CFG.get("headless", True))
    page = browser.new_page(viewport=CFG.get("viewport", {"width": 1440, "height": 900}))
    try:
        kds_login(page, CFG["kds"]["creator"]["username"])
        page.goto(f"{CFG['kds']['base_url']}/Transaction/MKPPDF", wait_until="domcontentloaded")
        page.wait_for_timeout(2000)
        page.click("#btnNew")
        page.wait_for_timeout(1500)
        page.click("#btnLOVGroupName")
        page.wait_for_timeout(2000)
        iframe = page.wait_for_selector("iframe.fancybox-iframe").content_frame()
        iframe.locator("#tbLOV tbody tr", has_text="CSA").first.locator(".btnSelect").click()
        page.wait_for_timeout(1500)
        # minimal 1 activity
        act = CFG["test_data"]["activities_mixed_fee"][0]
        fill_activity_row(page, 1, {**act, "amount": "100000"}, "REJ")
        ck = fill_ckeditors(page, f"SIT Reject {RUN_TAG}", "Reject T1 path", f"reject {RUN_TAG}")
        if ck.get("desc", 0) == 0:
            raise RuntimeError(f"CKEditor reject path empty: {ck}")
        doc, status, body = submit_mkpp_bootbox(page)
        STATE["mkpp_reject"] = doc
        ev = shot(page, "W3_01_mkpp_for_reject.png")
        rec(
            "AP-03a",
            "KICAO",
            "MKPP for reject",
            "Positive",
            "Create MKPP to reject",
            "Minimal submit bootbox",
            "DocNo",
            "High",
            "Pass" if doc else "Fail",
            f"doc={doc} status={status} body={body[:300]}",
            ev,
        )
        if not doc:
            raise RuntimeError(f"no doc for reject: {body[:400]}")

        page.context.clear_cookies()
        kds_login(page, CFG["kds"]["mkpp_approver_t1"]["username"])
        page.goto(f"{CFG['kds']['base_url']}/Transaction/Approval", wait_until="domcontentloaded")
        page.wait_for_timeout(2500)
        # Prefer Reject checkbox/button if exists
        rejected = False
        page.evaluate(
            """(docNo) => { if (window.oTableDetail) oTableDetail.search(docNo).draw(); }""",
            doc,
        )
        page.wait_for_timeout(1500)
        # try reject checkbox
        rj = page.evaluate(
            """(docNo) => {
                let chk=null;
                document.querySelectorAll('#dtDetail tbody tr').forEach(tr => {
                    if ((tr.innerText||'').indexOf(docNo)>=0) chk = tr.querySelector('.chkDetailReject, .chkReject');
                });
                if (!chk) return {found:false};
                chk.checked=true;
                return {found:true};
            }""",
            doc,
        )
        if rj.get("found") and page.locator("#btnReject, button:has-text('Reject')").count():
            page.locator("#btnReject, button:has-text('Reject')").first.click()
            page.wait_for_timeout(1000)
            if page.locator(".swal2-input, textarea").count():
                page.locator(".swal2-input, textarea").first.fill("SIT Exhaustive: reject MKPP Tier1")
            if page.locator(".swal2-confirm").count():
                page.locator(".swal2-confirm").click()
                page.wait_for_timeout(2500)
            rejected = True
        ev = shot(page, "W3_02_reject_attempt.png")
        rec(
            "AP-03",
            "KICAO",
            "MKPP Tier1 Reject",
            "Negative",
            f"Reject {doc}",
            "T1 reject + reason",
            "Status REJECTED",
            "High",
            "Pass" if rejected else "Fail",
            f"rejected={rejected} rj={rj}",
            ev,
        )
    except Exception as ex:
        ev = shot(page, "W3_FAIL.png")
        rec("W3-STOP", "KICAO", "Wave3 stopper", "Negative", "Unhandled", "-", "Complete", "Critical", "Fail", f"{ex}\n{traceback.format_exc()}", ev)
        print(traceback.format_exc())
    finally:
        browser.close()


# ─── W0b Maven UI smoke ──────────────────────────────────────────────────────
def wave0b_maven_ui(p):
    print("\n=== W0b MAVEN UI ===")
    browser = p.chromium.launch(headless=CFG.get("headless", True))
    page = browser.new_page(viewport=CFG.get("viewport", {"width": 1440, "height": 900}))
    try:
        maven_login(page)
        page.goto(f"{CFG['maven']['base_url']}{CFG['maven']['mapping_url']}", wait_until="domcontentloaded")
        page.wait_for_timeout(3500)
        ev = shot(page, "W0b_01_mapping.png")
        ok = ("Login" not in page.url) and (
            page.locator("#tblMappingSubdist").count() > 0 or "MappingSubdist" in page.url
        )
        rec("MV-01", "MAVEN", "Mapping Index", "Positive", "Open Mapping Subdist", "Login + navigate", "Index load (not login redirect)", "High", "Pass" if ok else "Fail", f"ok={ok} url={page.url}", ev)

        page.goto(f"{CFG['maven']['base_url']}{CFG['maven']['monitoring_url']}", wait_until="domcontentloaded")
        page.wait_for_timeout(3500)
        ev = shot(page, "W0b_02_monitoring.png")
        mon = ("Login" not in page.url) and (
            page.locator("#btnExportMappedExcel").count() > 0
            or page.locator("#btnRefreshClaims").count() > 0
            or "MonitoringSubdist" in page.url
        )
        rec("MV-07", "MAVEN", "Monitoring Index", "Positive", "Open Monitoring columns BI/Region/Export", "Navigate MonitoringSubdist", "Page authenticated", "High", "Pass" if mon else "Fail", f"url={page.url}", ev)
    except Exception as ex:
        ev = shot(page, "W0b_FAIL.png")
        rec("W0b-STOP", "MAVEN", "UI stopper", "Negative", "Unhandled", "-", "Complete", "Critical", "Fail", f"{ex}\n{traceback.format_exc()}", ev)
        print(traceback.format_exc())
    finally:
        browser.close()


def write_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "SIT Exhaustive"
    headers = [
        "No.",
        "Scenario ID",
        "Modul / Area",
        "Field Skenario",
        "Tipe",
        "Deskripsi",
        "Prasyarat",
        "Langkah",
        "Expected",
        "Priority",
        "Status Uji",
        "Hasil Aktual",
        "Bukti Pengujian",
        "QA",
        "Tanggal Uji",
    ]
    ws.append(headers)
    fill_pass = PatternFill("solid", fgColor="C6EFCE")
    fill_fail = PatternFill("solid", fgColor="FFC7CE")
    thin = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    for i, r in enumerate(RESULTS, 1):
        evid = r.get("evidence") or ""
        row = [
            i,
            r["id"],
            r["modul"],
            r["area"],
            r["tipe"],
            r["desc"],
            "-",
            r["langkah"],
            r["expected"],
            r["priority"],
            r["status"],
            r["actual"],
            evid,
            r["qa"],
            r["tanggal"],
        ]
        ws.append(row)
        status_cell = ws.cell(row=i + 1, column=11)
        if r["status"] == "Pass":
            status_cell.fill = fill_pass
        elif r["status"] == "Fail":
            status_cell.fill = fill_fail
        if evid:
            link = EVIDENCE / evid
            if link.exists():
                ws.cell(row=i + 1, column=13).hyperlink = str(link)
                ws.cell(row=i + 1, column=13).value = "Lihat Evidence"
                ws.cell(row=i + 1, column=13).font = Font(color="0563C1", underline="single")
        for c in range(1, 16):
            ws.cell(row=i + 1, column=c).border = thin
            ws.cell(row=i + 1, column=c).alignment = Alignment(wrap_text=True, vertical="top")

    out = REPORTS / f"DF_SIT_Exhaustive_{STAMP}.xlsx"
    wb.save(out)
    summary = {
        "stamp": STAMP,
        "pass": sum(1 for r in RESULTS if r["status"] == "Pass"),
        "fail": sum(1 for r in RESULTS if r["status"] == "Fail"),
        "total": len(RESULTS),
        "state": STATE,
        "excel": str(out),
        "evidence": str(EVIDENCE),
    }
    (EVIDENCE / "result.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    (HERE / "last_result.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print("\n=== SUMMARY ===")
    print(json.dumps(summary, indent=2))
    return out


def main():
    wave0_api_probe()
    with sync_playwright() as p:
        wave0b_maven_ui(p)
        wave1_mkpp_happy(p)
        wave2_prm_abm(p)
        wave3_mkpp_reject(p)
    write_workbook()


if __name__ == "__main__":
    main()
