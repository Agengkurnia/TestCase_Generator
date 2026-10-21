"""
Batch-run expanded MKPP presets (each total <= 100_000) through:
  Muslih create/submit -> PATRICK.YOHANNES approve
Optional: skip PRM/ABM (use --full for one small claim path).

Usage:
  py -3 run_expanded_batch.py
  py -3 run_expanded_batch.py --presets P-SMALL-10,P-OFFER-30,P-ADS-40
  py -3 run_expanded_batch.py --presets P-SMALL-10 --full
"""
from __future__ import annotations

import argparse
import json
import time
import traceback
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from playwright.sync_api import sync_playwright

from _pack import EVIDENCE_ROOT, REPORTS, ROOT
import run_e2e_corrected as e2e
from cases_expanded import MKPP_PRESETS, get_preset, validate_preset, BUDGET_CAP

HERE = ROOT
STAMP = datetime.now().strftime("%Y_%m_%d_%H-%M-%S")
EVIDENCE = EVIDENCE_ROOT / f"evidence_expanded_{STAMP}"
EVIDENCE.mkdir(parents=True, exist_ok=True)
RESULTS = []


def rec(sid, modul, status, actual, evidence=""):
    RESULTS.append(
        {
            "id": sid,
            "modul": modul,
            "status": status,
            "actual": str(actual)[:2000],
            "evidence": evidence,
            "qa": e2e.QA,
            "tanggal": e2e.TODAY,
        }
    )
    print(f"  [{'PASS' if status == 'Pass' else 'FAIL'}] {sid}: {str(actual)[:140]}")


def shot(page, name):
    path = EVIDENCE / name
    page.screenshot(path=str(path), full_page=False)
    return path.name


def create_and_approve(p, preset: dict) -> str:
    total = validate_preset(preset, BUDGET_CAP)
    pid = preset["id"]
    print(f"\n=== PRESET {pid} total={total} ===")
    browser = p.chromium.launch(headless=e2e.CFG.get("headless", True))
    page = browser.new_page(viewport=e2e.CFG.get("viewport", {"width": 1440, "height": 900}))
    doc = ""
    try:
        e2e.kds_login(page, e2e.KDS_CREATOR)
        page.goto(f"{e2e.KDS}/Transaction/MKPPDF", wait_until="domcontentloaded")
        page.wait_for_timeout(2500)
        page.click("#btnNew")
        page.wait_for_timeout(2000)
        page.click("#btnLOVGroupName")
        page.wait_for_timeout(2000)
        iframe = page.wait_for_selector("iframe.fancybox-iframe").content_frame()
        iframe.wait_for_load_state("networkidle")
        iframe.locator("#tbLOV tbody tr", has_text=preset["group"]).first.locator(".btnSelect").click()
        page.wait_for_timeout(2000)

        for i, a in enumerate(preset["activities"]):
            prefix = "BDO" if "BANDUNG" in a["outlet"].upper() else "JKT"
            # reuse add_activity (always 3 SKP); for skp_per_act=1 still OK (sum splits)
            e2e.add_activity(page, i + 1, a["outlet"], a["act_code"], a["amount"], prefix)
            rec(f"{pid}-ACT-{i+1}", "MKPP", "Pass", f"{a['act_code']}={a['amount']}", shot(page, f"{pid}_act_{i+1}.png"))

        ck = e2e.fill_ckeditors(
            page,
            f"Expanded {pid} {e2e.RUN_TAG}",
            f"Batch preset {pid} total={total}",
            f"QA Ageng {e2e.RUN_TAG}",
        )
        if ck.get("desc", 0) == 0:
            raise RuntimeError(f"CKEditor empty: {ck}")

        page.evaluate("window.scrollTo(0,0)")
        page.click("#btnSubmit")
        page.wait_for_timeout(1000)
        confirm = page.wait_for_selector("button[data-bb-handler=confirm]", timeout=15000)
        with page.expect_response(lambda r: "MKPPDF/SubmitData" in r.url, timeout=180000) as w:
            confirm.click()
        body = ""
        try:
            body = w.value.text()[:400]
        except Exception:
            pass
        try:
            ok = page.wait_for_selector("div.bootbox-alert button[data-bb-handler=ok]", timeout=20000)
            ok.click()
            page.wait_for_timeout(2000)
        except Exception:
            pass
        doc = page.input_value("#txtDocNo")
        if not doc:
            page.click("#btnFind")
            page.wait_for_timeout(2500)
            fr = page.wait_for_selector("iframe.fancybox-iframe").content_frame()
            fr.click("table tbody tr:first-child .btnSelect")
            page.wait_for_timeout(2000)
            doc = page.input_value("#txtDocNo")
        rec(f"{pid}-SUBMIT", "MKPP", "Pass" if doc else "Fail", f"doc={doc} body={body[:180]}", shot(page, f"{pid}_submit.png"))
        if not doc:
            raise RuntimeError("No DocNo")

        page.context.clear_cookies()
        e2e.kds_login(page, e2e.KDS_PATRICK)
        page.goto(f"{e2e.KDS}/Transaction/ApprovalDoc/Index", wait_until="domcontentloaded")
        page.wait_for_timeout(3000)
        page.click("#btnNew")
        page.wait_for_timeout(3000)
        picked = e2e.approve_only_target(page, doc)
        if not picked.get("found"):
            rec(f"{pid}-APPR", "Approval", "Fail", picked, shot(page, f"{pid}_patrick_miss.png"))
            raise RuntimeError(f"Patrick miss: {picked}")
        page.click("#btnSave")
        page.wait_for_timeout(1500)
        conf = page.query_selector('.bootbox button[data-bb-handler="confirm"]')
        if conf:
            conf.click()
            page.wait_for_timeout(5000)
        okb = page.query_selector('.bootbox button[data-bb-handler="ok"]')
        if okb:
            okb.click()
            page.wait_for_timeout(2000)
        rec(f"{pid}-APPR", "Approval", "Pass", doc, shot(page, f"{pid}_patrick_ok.png"))
        return doc
    except Exception as ex:
        rec(f"{pid}-STOP", "MKPP", "Fail", f"{ex}\n{traceback.format_exc()}", shot(page, f"{pid}_FAIL.png"))
        raise
    finally:
        browser.close()


def write_xlsx(docs):
    wb = Workbook()
    ws = wb.active
    ws.title = "Expanded Batch"
    ws.append(["No", "ID", "Modul", "Status", "Actual", "Evidence", "QA", "Tanggal"])
    fp, ff = PatternFill("solid", fgColor="C6EFCE"), PatternFill("solid", fgColor="FFC7CE")
    for i, r in enumerate(RESULTS, 1):
        ws.append([i, r["id"], r["modul"], r["status"], r["actual"], r["evidence"], r["qa"], r["tanggal"]])
        ws.cell(i + 1, 4).fill = fp if r["status"] == "Pass" else ff
        if r["evidence"] and (EVIDENCE / r["evidence"]).exists():
            c = ws.cell(i + 1, 6)
            c.value = "Lihat Evidence"
            c.hyperlink = str(EVIDENCE / r["evidence"])
            c.font = Font(color="0563C1", underline="single")
    out = REPORTS / f"DF_SIT_Expanded_Batch_{STAMP}.xlsx"
    wb.save(out)
    summary = {
        "pass": sum(1 for r in RESULTS if r["status"] == "Pass"),
        "fail": sum(1 for r in RESULTS if r["status"] == "Fail"),
        "total": len(RESULTS),
        "docs": docs,
        "cap": BUDGET_CAP,
        "excel": str(out),
        "evidence": str(EVIDENCE),
    }
    (EVIDENCE / "result.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    (HERE / "last_expanded_batch.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print("\n=== SUMMARY ===")
    print(json.dumps(summary, indent=2))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument(
        "--presets",
        default="P-SMALL-10,P-OFFER-30,P-ADS-40",
        help="Comma-separated preset ids (each <=100k)",
    )
    ap.add_argument("--full", action="store_true", help="Also run PRM+ABM for first preset doc")
    args = ap.parse_args()
    ids = [x.strip() for x in args.presets.split(",") if x.strip()]
    presets = [get_preset(i) for i in ids]
    for p in presets:
        validate_preset(p)

    docs = []
    with sync_playwright() as p:
        for preset in presets:
            try:
                doc = create_and_approve(p, preset)
                docs.append({"preset": preset["id"], "doc": doc})
                time.sleep(1)
            except Exception:
                docs.append({"preset": preset["id"], "doc": "", "error": True})
                continue

        if args.full and docs and docs[0].get("doc"):
            e2e.RESULTS.clear()
            e2e.STATE["mkpp"] = docs[0]["doc"]
            e2e.STATE["claims"] = []
            e2e.EVIDENCE = EVIDENCE
            try:
                e2e.step_prm_claims(p, docs[0]["doc"])
                e2e.step_abm_muslih(p)
                # merge e2e results into batch RESULTS
                for r in e2e.RESULTS:
                    RESULTS.append(r)
            except Exception as ex:
                rec("FULL-STOP", "PRM/ABM", "Fail", str(ex))

    write_xlsx(docs)


if __name__ == "__main__":
    main()
