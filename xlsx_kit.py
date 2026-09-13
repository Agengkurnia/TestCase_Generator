"""Shared Excel styling + helpers for all product testcase packs."""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

GREEN = PatternFill("solid", fgColor="1E7D37")
GREEN_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
TITLE_FONT = Font(bold=True, name="Calibri", size=14, color="1E7D37")
CELL_FONT = Font(name="Calibri", size=11)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(
    left=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
    top=Side(style="thin", color="D0D0D0"),
    bottom=Side(style="thin", color="D0D0D0"),
)
ALT = PatternFill("solid", fgColor="F4F7F4")

SKN_HEADERS = [
    "No.",
    "Scenario ID",
    "Modul / Area",
    "Field Skenario",
    "Tipe",
    "Deskripsi Skenario",
    "Prasyarat",
    "Langkah Uji",
    "Hasil yang Diharapkan",
    "Priority",
    "Status Uji",
    "Hasil Aktual",
    "QA",
    "Tanggal Uji",
]

SIT_HEADERS = [
    "No.",
    "Issue ID",
    "Ref ID",
    "Modul / Area",
    "Field Skenario",
    "Deskripsi Skenario",
    "Deskripsi Issue / Bug",
    "Langkah Reproduksi",
    "Gambar",
    "Priority / Severity",
    "Status Issue",
    "Assignee / Dev",
    "Tanggal Ditemukan",
    "Tanggal Selesai",
]


def style_header(ws, row, ncols):
    for col in range(1, ncols + 1):
        cell = ws.cell(row, col)
        cell.fill = GREEN
        cell.font = GREEN_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = THIN


def write_rows(ws, start_row, rows):
    for i, row in enumerate(rows):
        r = start_row + i
        fill = ALT if i % 2 else None
        for c, val in enumerate(row, 1):
            cell = ws.cell(r, c, val)
            cell.font = CELL_FONT
            cell.alignment = WRAP
            cell.border = THIN
            if fill:
                cell.fill = fill


def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
